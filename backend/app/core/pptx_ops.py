"""Операции над пакетом PPTX: клонирование слайдов, удаление, работа с текстом.

Все слайды итоговой колоды — копии слайдов шаблона (с его макетами и мастерами),
поэтому аудит «слайд собран на макете шаблона» выполняется по построению.
"""
from __future__ import annotations

import copy
import re
from typing import Iterable, Optional

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.slide import Slide

from ..parsing.ooxml import NS, first, qn, xp

R_ATTRS = (qn("r:embed"), qn("r:link"), qn("r:id"), qn("r:pict"))
SKIP_RELS = (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE)


def top_shapes(slide: Slide) -> list:
    tree = slide.shapes._spTree
    return [c for c in tree if etree.QName(c).localname in ("sp", "pic", "grpSp", "graphicFrame", "cxnSp", "contentPart")]


def shape_id(el) -> Optional[str]:
    nv = first(el, "*[1]/p:cNvPr")
    return nv.get("id") if nv is not None else None


def clone_slide(prs: Presentation, src: Slide) -> Slide:
    """Копия слайда src в конец презентации prs (тот же пакет)."""
    new = prs.slides.add_slide(src.slide_layout)
    new_tree = new.shapes._spTree
    for el in list(new_tree):
        if etree.QName(el).localname in ("sp", "pic", "grpSp", "graphicFrame", "cxnSp", "contentPart"):
            new_tree.remove(el)
    rid_map: dict[str, str] = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype in SKIP_RELS:
            continue
        if rel.is_external:
            rid_map[rid] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        else:
            rid_map[rid] = new.part.relate_to(rel.target_part, rel.reltype)
    for el in top_shapes(src):
        c = copy.deepcopy(el)
        _remap_rids(c, rid_map)
        new_tree.append(c)
    # фон слайда
    src_bg = first(src._element, "p:cSld/p:bg")
    if src_bg is not None:
        cs = first(new._element, "p:cSld")
        old = first(cs, "p:bg")
        if old is not None:
            cs.remove(old)
        bg = copy.deepcopy(src_bg)
        _remap_rids(bg, rid_map)
        cs.insert(0, bg)
    # переопределение цветовой схемы
    src_ovr = first(src._element, "p:clrMapOvr")
    if src_ovr is not None:
        old = first(new._element, "p:clrMapOvr")
        if old is not None:
            new._element.remove(old)
        new._element.append(copy.deepcopy(src_ovr))
    return new


def _remap_rids(el, rid_map: dict[str, str]):
    for node in el.iter():
        for a in R_ATTRS:
            v = node.get(a)
            if v and v in rid_map:
                node.set(a, rid_map[v])


def delete_slides(prs: Presentation, keep: Iterable[Slide]) -> None:
    keep_ids = {id(s._element) for s in keep}
    sld_lst = prs.slides._sldIdLst
    for sld_id in list(sld_lst):
        rid = sld_id.get(qn("r:id"))
        slide = prs.part.related_part(rid).slide
        if id(slide._element) not in keep_ids:
            sld_lst.remove(sld_id)
            prs.part.drop_rel(rid)
    # секции ссылаются на id слайдов — удаляем, чтобы не повредить файл
    for ext in xp(prs.part._element, "p:extLst/p:ext"):
        if first(ext, "p14:sectionLst") is not None:
            ext.getparent().remove(ext)
    cust = first(prs.part._element, "p:custShowLst")
    if cust is not None:
        prs.part._element.remove(cust)


def reorder_slides(prs: Presentation, order: list[Slide]) -> None:
    sld_lst = prs.slides._sldIdLst
    by_el = {}
    for sld_id in sld_lst:
        s = prs.part.related_part(sld_id.get(qn("r:id"))).slide
        by_el[id(s._element)] = sld_id
    for s in order:
        el = by_el[id(s._element)]
        sld_lst.remove(el)
        sld_lst.append(el)


def remove_shapes(slide: Slide, sids: set[str]) -> None:
    tree = slide.shapes._spTree
    for el in top_shapes(slide):
        if shape_id(el) in sids:
            tree.remove(el)


def find_shape(slide: Slide, sid: str):
    for el in slide.shapes._spTree.iter():
        if etree.QName(el).localname in ("sp", "pic", "grpSp", "graphicFrame", "cxnSp"):
            if shape_id(el) == sid:
                return el
    return None


def next_shape_id(slide: Slide) -> int:
    ids = [int(v) for v in xp(slide.shapes._spTree, ".//p:cNvPr/@id") if str(v).isdigit()]
    return max(ids + [1]) + 1


# ------------------------------------------------------------------ текст

def _template_rpr(sp_el):
    """rPr первого непустого прогона (чтобы сохранить форматирование шаблона)."""
    for r in xp(sp_el, ".//a:p/a:r"):
        t = first(r, "a:t")
        if t is not None and (t.text or "").strip():
            rpr = first(r, "a:rPr")
            return copy.deepcopy(rpr) if rpr is not None else None
    e = first(sp_el, ".//a:p/a:endParaRPr")
    if e is not None:
        e = copy.deepcopy(e)
        e.tag = qn("a:rPr")
        return e
    return None


def _template_ppr(sp_el):
    for p in xp(sp_el, ".//a:txBody/a:p"):
        if "".join(p.itertext()).strip():
            ppr = first(p, "a:pPr")
            return copy.deepcopy(ppr) if ppr is not None else None
    p = first(sp_el, ".//a:txBody/a:p/a:pPr")
    return copy.deepcopy(p) if p is not None else None


def set_text(sp_el, paragraphs: list, size_pt: Optional[float] = None, color: Optional[str] = None,
             bold: Optional[bool] = None, keep_breaks: bool = False, font: Optional[str] = None) -> None:
    """Заменяет текст фигуры, сохраняя rPr/pPr первого абзаца шаблона.

    paragraphs: список строк или словарей {"text", "bold", "color", "size", "bullet": bool}.
    """
    txb = first(sp_el, "p:txBody")
    if txb is None:
        return
    rpr_t = _template_rpr(sp_el)
    ppr_t = _template_ppr(sp_el)
    for p in xp(txb, "a:p"):
        txb.remove(p)
    for item in paragraphs:
        if isinstance(item, str):
            item = {"text": item}
        p = etree.SubElement(txb, qn("a:p"))
        if ppr_t is not None:
            p.append(copy.deepcopy(ppr_t))
        if item.get("bullet") is False:
            ppr = first(p, "a:pPr")
            if ppr is not None:
                for b in xp(ppr, "a:buChar | a:buAutoNum | a:buFont | a:buSzPct"):
                    ppr.remove(b)
        text = str(item.get("text", ""))
        chunks = text.split("\n") if not keep_breaks else [text]
        for ci, chunk in enumerate(chunks):
            if ci > 0:
                br = etree.SubElement(p, qn("a:br"))
                if rpr_t is not None:
                    br.append(copy.deepcopy(rpr_t))
            r = etree.SubElement(p, qn("a:r"))
            rpr = copy.deepcopy(rpr_t) if rpr_t is not None else etree.Element(qn("a:rPr"))
            rpr.set("lang", "ru-RU")
            if rpr.get("dirty") is not None:
                del rpr.attrib["dirty"]
            sz = item.get("size", size_pt)
            if sz:
                rpr.set("sz", str(int(round(sz * 100))))
            b = item.get("bold", bold)
            if b is not None:
                rpr.set("b", "1" if b else "0")
            col = item.get("color", color)
            if col:
                set_run_color(rpr, col)
            if font:
                set_run_font(rpr, font)
            r.append(rpr)
            t = etree.SubElement(r, qn("a:t"))
            t.text = chunk
            if chunk != chunk.strip():
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def set_run_color(rpr, hexv: str) -> None:
    for old in xp(rpr, "a:solidFill | a:gradFill | a:noFill"):
        rpr.remove(old)
    sf = etree.Element(qn("a:solidFill"))
    c = etree.SubElement(sf, qn("a:srgbClr"))
    c.set("val", hexv.upper())
    # solidFill должен идти после ln и до latin/ea/cs
    idx = 0
    for i, ch in enumerate(rpr):
        if etree.QName(ch).localname == "ln":
            idx = i + 1
    rpr.insert(idx, sf)


def set_run_font(rpr, family: str) -> None:
    """Гарнитура для latin/cs (порядок детей rPr: ln, fill, effect, highlight, uLn, latin, ea, cs, sym...)."""
    for tag in ("a:latin", "a:cs"):
        old = first(rpr, tag)
        if old is not None:
            old.set("typeface", family)
            continue
        el = etree.Element(qn(tag))
        el.set("typeface", family)
        after = [c for c in rpr if etree.QName(c).localname in ("ln", "solidFill", "gradFill", "noFill", "effectLst",
                                                                  "highlight", "uLnTx", "uLn", "uFillTx", "uFill", "latin", "ea")]
        if after:
            after[-1].addnext(el)
        else:
            rpr.insert(0, el)


def set_xfrm(el, x: int, y: int, w: int, h: int) -> None:
    tag = etree.QName(el).localname
    if tag == "graphicFrame":
        xf = first(el, "p:xfrm")
    elif tag == "grpSp":
        xf = first(el, "p:grpSpPr/a:xfrm")
    else:
        xf = first(el, "p:spPr/a:xfrm")
    if xf is None:
        sppr = first(el, "p:spPr")
        if sppr is None:
            return
        xf = etree.Element(qn("a:xfrm"))
        sppr.insert(0, xf)
    off = first(xf, "a:off")
    ext = first(xf, "a:ext")
    if off is None:
        off = etree.SubElement(xf, qn("a:off"))
    if ext is None:
        ext = etree.SubElement(xf, qn("a:ext"))
    if tag == "grpSp":
        # пересчёт дочерних координат не нужен: меняем только внешнюю рамку
        pass
    off.set("x", str(int(x)))
    off.set("y", str(int(y)))
    ext.set("cx", str(int(max(w, 1))))
    ext.set("cy", str(int(max(h, 1))))


def set_autofit_none(sp_el) -> None:
    bp = first(sp_el, "p:txBody/a:bodyPr")
    if bp is None:
        return
    for old in xp(bp, "a:normAutofit | a:spAutoFit | a:noAutofit"):
        bp.remove(old)
    warp = first(bp, "a:prstTxWarp")
    bp.insert(bp.index(warp) + 1 if warp is not None else 0, etree.Element(qn("a:noAutofit")))


def clean_text(s: str) -> str:
    s = re.sub(r"\s+", " ", s or "").strip()
    return s
