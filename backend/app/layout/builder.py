"""Сборка PPTX из SlideLayout: клонирование холстов шаблона + нативные объекты."""
from __future__ import annotations

import copy
import io
from pathlib import Path
from typing import Optional
from xml.sax.saxutils import escape

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor

from ..core import pptx_ops
from ..parsing.model import Box, CardStyle, TemplateProfile
from ..parsing.ooxml import NS, first, qn, xp
from . import icons
from .ir import El, SlideLayout

NSDECL = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
          'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
          'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')
NO_STYLE_TABLE = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"


def _xml(s: str):
    return etree.fromstring(s)


def _solid(hexv: Optional[str]) -> str:
    return f'<a:solidFill><a:srgbClr val="{hexv}"/></a:solidFill>' if hexv else "<a:noFill/>"


class DeckBuilder:
    def __init__(self, profile: TemplateProfile):
        self.p = profile
        self.prs = Presentation(profile.file)
        self.src = list(self.prs.slides)
        self.made = []
        self.W, self.H = profile.slide_w, profile.slide_h

    # ================================================================ slide
    def add(self, lay: SlideLayout):
        canvas = self.p.canvas(lay.canvas)
        src = self.src[canvas.source_slide - 1]
        slide = pptx_ops.clone_slide(self.prs, src)
        keep = set(canvas.keep_sids) - set(lay.remove_sids)
        all_top = {pptx_ops.shape_id(el) for el in pptx_ops.top_shapes(slide)}
        pptx_ops.remove_shapes(slide, all_top - keep)
        self._slots(slide, lay, canvas)
        self._next_id = pptx_ops.next_shape_id(slide)
        for el in lay.elements:
            try:
                self._element(slide, el)
            except Exception as e:  # элемент не должен ронять всю колоду
                lay.warnings.append(f"element_failed:{el.kind}:{e.__class__.__name__}")
        if lay.notes:
            try:
                slide.notes_slide.notes_text_frame.text = lay.notes
            except Exception:
                pass
        self.made.append(slide)
        return slide

    def save(self, path: Path) -> Path:
        pptx_ops.delete_slides(self.prs, self.made)
        self.prs.save(str(path))
        return path

    # ================================================================ slots
    def _slots(self, slide, lay: SlideLayout, canvas):
        by_sid = {s.sid: s for s in canvas.slots}
        for sf in lay.slots:
            el = pptx_ops.find_shape(slide, sf.sid)
            if el is None:
                continue
            paras = sf.paras or [sf.text]
            pptx_ops.set_text(el, paras, size_pt=sf.size, color=sf.color, font=sf.font)
            pptx_ops.set_autofit_none(el)
            slot = by_sid.get(sf.sid)
            if slot is not None and (sf.h or sf.w):
                pptx_ops.set_xfrm(el, slot.box.x, slot.box.y, sf.w or slot.box.w, sf.h or slot.box.h)
        title = next((s for s in canvas.slots if s.role == "title"), None)
        if lay.label_width and canvas.title_label_sid and title is not None:
            lab = pptx_ops.find_shape(slide, canvas.title_label_sid)
            if lab is not None:
                xf = first(lab, "p:spPr/a:xfrm")
                if xf is not None:
                    off, ext = first(xf, "a:off"), first(xf, "a:ext")
                    x = int(off.get("x"))
                    text_left = title.box.x + 91440
                    pad = max(int(0.12 * INCH_EMU), text_left - x)
                    w = lay.label_width + 2 * pad
                    limit = (title.box.x + (canvas.title_max_w or self.W)) - x
                    w = max(int(0.6 * INCH_EMU), min(w, limit, int(0.92 * self.W) - x))
                    ext.set("cx", str(int(w)))

    # ================================================================ elements
    def _nid(self) -> int:
        self._next_id += 1
        return self._next_id

    def _append(self, slide, el_xml):
        slide.shapes._spTree.append(el_xml)

    def _element(self, slide, el: El):
        k = el.kind
        if k == "text":
            self._append(slide, self._text_xml(el))
        elif k == "card":
            self._append(slide, self._card_xml(slide, el))
        elif k == "shape":
            self._append(slide, self._shape_xml(el))
        elif k == "line":
            self._append(slide, self._line_xml(el))
        elif k == "icon":
            name = icons.find_icon(el.icon) or "circle-check"
            data = icons.render_icon(name, el.color or "000000")
            pic = slide.shapes.add_picture(io.BytesIO(data), Emu(el.box.x), Emu(el.box.y), Emu(el.box.w), Emu(el.box.h))
            pic.name = f"icon|{el.id}"
        elif k == "image":
            self._image(slide, el)
        elif k == "chart":
            self._chart(slide, el)
        elif k == "table":
            self._table(slide, el)

    # ---------------------------------------------------------------- text
    def _text_xml(self, el: El):
        b = el.box
        sz = int(round((el.size or 14) * 100))
        paras_xml = []
        anchor = {"t": "t", "ctr": "ctr", "b": "b"}.get(el.valign, "t")
        for p in el.paras:
            ppr = [f'<a:pPr algn="{el.align}"']
            if p.bullet:
                mar = int((el.size or 14) * 12700 * 1.1)
                ppr.append(f' marL="{mar}" indent="-{mar}"')
            ppr.append(">")
            ppr.append(f'<a:lnSpc><a:spcPct val="{int(el.line_spacing * 100000)}"/></a:lnSpc>')
            if p.space_after:
                ppr.append(f'<a:spcAft><a:spcPts val="{int(p.space_after * 100)}"/></a:spcAft>')
            if p.bullet:
                bc = p.bullet_color or el.color or "000000"
                ppr.append(f'<a:buClr><a:srgbClr val="{bc}"/></a:buClr><a:buFont typeface="Arial"/><a:buChar char="&#8226;"/>')
            else:
                ppr.append("<a:buNone/>")
            ppr.append("</a:pPr>")
            runs = []
            for r in p.runs:
                if not r.text:
                    continue
                bold = r.bold if r.bold is not None else el.bold
                col = r.color or el.color or "000000"
                cap = ' cap="all"' if el.caps else ""
                font = escape(el.font or "Arial", {'"': "&quot;"})
                t = escape(r.text)
                space = ' xml:space="preserve"' if t != t.strip() else ""
                runs.append(f'<a:r><a:rPr lang="ru-RU" sz="{sz}" b="{1 if bold else 0}"{cap} dirty="0">{_solid(col)}'
                            f'<a:latin typeface="{font}"/><a:cs typeface="{font}"/></a:rPr><a:t{space}>{t}</a:t></a:r>')
            end = f'<a:endParaRPr lang="ru-RU" sz="{sz}" dirty="0"/>'
            paras_xml.append(f"<a:p>{''.join(ppr)}{''.join(runs)}{end}</a:p>")
        nid = self._nid()
        xml = (f'<p:sp {NSDECL}><p:nvSpPr><p:cNvPr id="{nid}" name="{el.role or "text"}|{el.id}"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
               f'<p:spPr><a:xfrm><a:off x="{b.x}" y="{b.y}"/><a:ext cx="{max(1, b.w)}" cy="{max(1, b.h)}"/></a:xfrm>'
               f'<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>'
               f'<p:txBody><a:bodyPr wrap="square" lIns="0" tIns="0" rIns="0" bIns="0" anchor="{anchor}" rtlCol="0"><a:noAutofit/></a:bodyPr>'
               f'<a:lstStyle/>{"".join(paras_xml) or "<a:p/>"}</p:txBody></p:sp>')
        return _xml(xml)

    # ---------------------------------------------------------------- card
    def _card_xml(self, slide, el: El):
        b = el.box
        nid = self._nid()
        cs: Optional[CardStyle] = self.p.card_styles[el.card] if el.card is not None and el.card < len(self.p.card_styles) else None
        if cs is not None:
            sppr = _xml(cs.sppr_xml)
            xf = first(sppr, "a:xfrm")
            ow = oh = None
            if xf is not None:
                ext = first(xf, "a:ext")
                if ext is not None:
                    ow, oh = int(ext.get("cx", 0)), int(ext.get("cy", 0))
                sppr.remove(xf)
            xf = _xml(f'<a:xfrm {NSDECL}><a:off x="{b.x}" y="{b.y}"/><a:ext cx="{b.w}" cy="{b.h}"/></a:xfrm>')
            sppr.insert(0, xf)
            geom = first(sppr, "a:prstGeom")
            if geom is not None and geom.get("prst") in ("roundRect",) and ow and oh:
                gd = first(geom, "a:avLst/a:gd")
                adj = int(gd.get("fmla", "val 16667").split()[-1]) if gd is not None else 16667
                r_abs = adj / 100000 * min(ow, oh)
                # не даём скруглению «съесть» маленькие карточки
                r_abs = min(r_abs, 0.22 * min(b.w, b.h))
                new_adj = int(max(0, min(50000, r_abs / max(1, min(b.w, b.h)) * 100000)))
                av = first(geom, "a:avLst")
                for g in list(av):
                    av.remove(g)
                etree.SubElement(av, qn("a:gd"), name="adj", fmla=f"val {new_adj}")
            self._remap_blips(slide, sppr, cs.source_slide)
            sp = _xml(f'<p:sp {NSDECL}><p:nvSpPr><p:cNvPr id="{nid}" name="card|{el.id}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr></p:sp>')
            sp.append(sppr)
            if cs.style_xml:
                st = _xml(cs.style_xml)
                self._remap_blips(slide, st, cs.source_slide)
                sp.append(st)
            return sp
        radius = el.radius or int(0.08 * 914400 * self.W / (13.333 * 914400))
        adj = int(max(0, min(50000, radius / max(1, min(b.w, b.h)) * 100000)))
        line = f'<a:ln w="9525">{_solid(el.line)}</a:ln>' if el.line else "<a:ln><a:noFill/></a:ln>"
        return _xml(f'<p:sp {NSDECL}><p:nvSpPr><p:cNvPr id="{nid}" name="card|{el.id}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
                    f'<p:spPr><a:xfrm><a:off x="{b.x}" y="{b.y}"/><a:ext cx="{b.w}" cy="{b.h}"/></a:xfrm>'
                    f'<a:prstGeom prst="roundRect"><a:avLst><a:gd name="adj" fmla="val {adj}"/></a:avLst></a:prstGeom>'
                    f'{_solid(el.fill)}{line}</p:spPr></p:sp>')

    def _remap_blips(self, slide, node, source_slide: int):
        blips = [n for n in node.iter() if n.get(qn("r:embed"))]
        if not blips:
            return
        src = self.src[source_slide - 1]
        for bl in blips:
            rid = bl.get(qn("r:embed"))
            if rid in src.part.rels:
                new_rid = slide.part.relate_to(src.part.rels[rid].target_part, src.part.rels[rid].reltype)
                bl.set(qn("r:embed"), new_rid)

    # ---------------------------------------------------------------- shapes
    def _shape_xml(self, el: El):
        b = el.box
        nid = self._nid()
        line = f'<a:ln w="{int(el.line_w * 12700) or 9525}">{_solid(el.line)}</a:ln>' if el.line else "<a:ln><a:noFill/></a:ln>"
        av = "<a:avLst/>"
        if el.geom == "roundRect" and el.radius:
            adj = int(max(0, min(50000, el.radius / max(1, min(b.w, b.h)) * 100000)))
            av = f'<a:avLst><a:gd name="adj" fmla="val {adj}"/></a:avLst>'
        return _xml(f'<p:sp {NSDECL}><p:nvSpPr><p:cNvPr id="{nid}" name="{el.role or "shape"}|{el.id}"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
                    f'<p:spPr><a:xfrm><a:off x="{b.x}" y="{b.y}"/><a:ext cx="{max(1, b.w)}" cy="{max(1, b.h)}"/></a:xfrm>'
                    f'<a:prstGeom prst="{el.geom}">{av}</a:prstGeom>{_solid(el.fill)}{line}</p:spPr></p:sp>')

    def _line_xml(self, el: El):
        b = el.box
        nid = self._nid()
        w = int((el.line_w or 1) * 12700)
        tail = '<a:tailEnd type="triangle" w="med" len="med"/>' if el.arrow else ""
        return _xml(f'<p:cxnSp {NSDECL}><p:nvCxnSpPr><p:cNvPr id="{nid}" name="{el.role or "line"}|{el.id}"/><p:cNvCxnSpPr/><p:nvPr/></p:nvCxnSpPr>'
                    f'<p:spPr><a:xfrm><a:off x="{b.x}" y="{b.y}"/><a:ext cx="{max(0, b.w)}" cy="{max(0, b.h)}"/></a:xfrm>'
                    f'<a:prstGeom prst="line"><a:avLst/></a:prstGeom>'
                    f'<a:ln w="{w}" cap="rnd">{_solid(el.line or "000000")}<a:round/>{tail}</a:ln></p:spPr></p:cxnSp>')

    # ---------------------------------------------------------------- image
    def _image(self, slide, el: El):
        path = Path(el.image)
        if not path.exists():
            return
        b = el.box
        with Image.open(path) as im:
            iw, ih = im.size
        pic = slide.shapes.add_picture(str(path), Emu(b.x), Emu(b.y), Emu(b.w), Emu(b.h))
        # cover-обрезка без искажения пропорций
        box_ar, img_ar = b.w / max(1, b.h), iw / max(1, ih)
        if img_ar > box_ar:
            cut = (1 - box_ar / img_ar) / 2
            pic.crop_left = pic.crop_right = cut
        elif img_ar < box_ar:
            cut = (1 - img_ar / box_ar) / 2
            pic.crop_top = pic.crop_bottom = cut
        pic.name = f"image|{el.id}"
        if el.radius:
            geom = first(pic._element, "p:spPr/a:prstGeom")
            if geom is not None:
                geom.set("prst", "roundRect")
                av = first(geom, "a:avLst")
                if av is None:
                    av = etree.SubElement(geom, qn("a:avLst"))
                adj = int(max(0, min(50000, el.radius / max(1, min(b.w, b.h)) * 100000)))
                etree.SubElement(av, qn("a:gd"), name="adj", fmla=f"val {adj}")

    # ---------------------------------------------------------------- chart
    def _chart(self, slide, el: El):
        c = el.chart or {}
        b = el.box
        typ = c.get("type", "column")
        cd = CategoryChartData()
        cats = [str(x) for x in c.get("categories", [])]
        cd.categories = cats
        series = c.get("series", [])[:5]
        for s in series:
            vals = list(s.get("values", []))[: len(cats)]
            vals += [0] * (len(cats) - len(vals))
            cd.add_series(s.get("name") or "Ряд", vals)
        xl = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
              "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE, "doughnut": XL_CHART_TYPE.DOUGHNUT}[typ]
        gf = slide.shapes.add_chart(xl, Emu(b.x), Emu(b.y), Emu(b.w), Emu(b.h), cd)
        ch = gf.chart
        colors = c.get("colors") or ["0077FF"]
        font = c.get("font", "Arial")
        fs = c.get("font_size", 11)
        tcol = c.get("text_color", "333333")
        ch.font.name = font
        ch.font.size = Pt(fs)
        ch.font.color.rgb = RGBColor.from_string(tcol)
        ch.has_title = False
        plot = ch.plots[0]
        multi = len(series) > 1
        if typ in ("pie", "doughnut"):
            ser = plot.series[0]
            for j, _ in enumerate(cats):
                pt = ser.points[j]
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = RGBColor.from_string(colors[j % len(colors)])
                pt.format.line.color.rgb = RGBColor.from_string(c.get("bg", "FFFFFF"))
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.show_percentage = typ == "pie"
            dl.show_value = typ != "pie"
            dl.font.size = Pt(fs)
            dl.font.bold = True
            ch.has_legend = True
            ch.legend.position = XL_LEGEND_POSITION.RIGHT if b.w > b.h * 1.3 else XL_LEGEND_POSITION.BOTTOM
            ch.legend.include_in_layout = False
            ch.legend.font.size = Pt(fs)
            if typ == "doughnut":
                hole = first(ch._chartSpace, ".//c:doughnutChart/c:holeSize")
                if hole is not None:
                    hole.set("val", "58")
        else:
            for i, ser in enumerate(plot.series):
                col = RGBColor.from_string(colors[i % len(colors)])
                if typ == "line":
                    ser.format.line.color.rgb = col
                    ser.format.line.width = Pt(2.25)
                    ser.smooth = False
                    try:
                        ser.marker.format.fill.solid()
                        ser.marker.format.fill.fore_color.rgb = col
                        ser.marker.format.line.color.rgb = col
                    except Exception:
                        pass
                else:
                    ser.format.fill.solid()
                    ser.format.fill.fore_color.rgb = col
                    ser.format.line.fill.background()
            if typ in ("column", "bar"):
                plot.gap_width = 70 if not multi else 90
                plot.overlap = -10 if multi else 0
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.font.size = Pt(max(7, fs - 1))
            dl.font.color.rgb = RGBColor.from_string(c.get("label_color", tcol))
            dl.number_format = "General"
            dl.number_format_is_linked = False
            try:
                dl.position = XL_LABEL_POSITION.OUTSIDE_END if typ != "line" else XL_LABEL_POSITION.ABOVE
            except Exception:
                pass
            va = ch.value_axis
            va.has_major_gridlines = True
            va.major_gridlines.format.line.color.rgb = RGBColor.from_string(c.get("grid_color", "DDDDDD"))
            va.major_gridlines.format.line.width = Pt(0.5)
            va.format.line.fill.background()
            va.tick_labels.font.size = Pt(max(7, fs - 1))
            va.tick_labels.font.color.rgb = RGBColor.from_string(tcol)
            if c.get("unit"):
                va.has_title = True
                va.axis_title.text_frame.text = c["unit"]
                r = va.axis_title.text_frame.paragraphs[0].runs[0]
                r.font.size = Pt(max(7, fs - 1))
                r.font.bold = False
                r.font.color.rgb = RGBColor.from_string(tcol)
                r.font.name = font
            ca = ch.category_axis
            ca.tick_labels.font.size = Pt(fs)
            ca.tick_labels.font.color.rgb = RGBColor.from_string(tcol)
            ca.format.line.color.rgb = RGBColor.from_string(c.get("grid_color", "DDDDDD"))
            ca.has_major_gridlines = False
            ch.has_legend = multi
            if multi:
                ch.legend.position = XL_LEGEND_POSITION.TOP
                ch.legend.include_in_layout = False
                ch.legend.font.size = Pt(fs)
        # прозрачный фон области диаграммы
        cs = ch._chartSpace
        for parent_path in ("c:chart/c:plotArea", "."):
            parent = first(cs, parent_path) if parent_path != "." else cs
            if parent is None:
                continue
            sppr = first(parent, "c:spPr")
            if sppr is None:
                sppr = etree.SubElement(parent, qn("c:spPr"))
                # c:spPr в chartSpace должен идти до c:txPr
                txpr = first(parent, "c:txPr")
                if txpr is not None and parent is cs:
                    parent.remove(sppr)
                    txpr.addprevious(sppr)
                elif parent is not cs:
                    # в plotArea spPr — после осей, перед extLst
                    ext = first(parent, "c:extLst")
                    if ext is not None:
                        parent.remove(sppr)
                        ext.addprevious(sppr)
            for ch_ in list(sppr):
                sppr.remove(ch_)
            etree.SubElement(sppr, qn("a:noFill"))
            ln = etree.SubElement(sppr, qn("a:ln"))
            etree.SubElement(ln, qn("a:noFill"))
        gf.name = f"chart|{el.id}"

    # ---------------------------------------------------------------- table
    def _table(self, slide, el: El):
        t = el.table or {}
        cols, rows = t["columns"], t["rows"]
        b = el.box
        nrows, ncols = len(rows) + 1, len(cols)
        gf = slide.shapes.add_table(nrows, ncols, Emu(b.x), Emu(b.y), Emu(b.w), Emu(b.h))
        gf.name = f"table|{el.id}"
        tbl = gf.table
        tbl_el = gf._element.graphic.graphicData.tbl
        tblpr = first(tbl_el, "a:tblPr")
        for a in ("firstRow", "bandRow", "firstCol", "lastRow", "lastCol", "bandCol"):
            if tblpr.get(a) is not None:
                del tblpr.attrib[a]
        sid = first(tblpr, "a:tableStyleId")
        if sid is None:
            sid = etree.SubElement(tblpr, qn("a:tableStyleId"))
        sid.text = NO_STYLE_TABLE
        for j, w in enumerate(t["widths"]):
            tbl.columns[j].width = Emu(int(w))
        for i, h in enumerate(t["heights"]):
            tbl.rows[i].height = Emu(int(h))
        size = t.get("size", 11)
        font = t.get("font", "Arial")
        pad = int(t.get("pad", 91440))
        border = t.get("border", "CCCCCC")
        for i in range(nrows):
            data = cols if i == 0 else rows[i - 1]
            for j in range(ncols):
                cell = tbl.cell(i, j)
                txt = str(data[j]) if j < len(data) else ""
                header = i == 0
                if header:
                    fill, color, bold = t["header_fill"], t["header_text"], True
                else:
                    fill = t.get("alt_fill") if (i % 2 == 0 and t.get("alt_fill")) else t.get("body_fill")
                    color, bold = t["body_text"], bool(t.get("first_col_bold") and j == 0)
                tf_ = cell.text_frame
                tf_.word_wrap = True
                cell.margin_left = cell.margin_right = Emu(pad)
                cell.margin_top = cell.margin_bottom = Emu(pad)
                p = tf_.paragraphs[0]
                p.text = ""
                r = p.add_run()
                r.text = txt
                r.font.size = Pt(size)
                r.font.bold = bold
                r.font.name = font
                r.font.color.rgb = RGBColor.from_string(color)
                tc = cell._tc
                tcpr = tc.get_or_add_tcPr()
                for ch_ in list(tcpr):
                    tcpr.remove(ch_)
                bw = 19050 if header else 9525
                for side in ("lnL", "lnR", "lnT", "lnB"):
                    if side == "lnB" or (side == "lnT" and i == 0):
                        ln = _xml(f'<a:{side} {NSDECL} w="{bw if side == "lnB" else 0}"><a:solidFill><a:srgbClr val="{border}"/></a:solidFill></a:{side}>'
                                  if side == "lnB" else f'<a:{side} {NSDECL} w="0"><a:noFill/></a:{side}>')
                    else:
                        ln = _xml(f'<a:{side} {NSDECL} w="0"><a:noFill/></a:{side}>')
                    tcpr.append(ln)
                tcpr.append(_xml(f'<a:solidFill {NSDECL}><a:srgbClr val="{fill}"/></a:solidFill>') if fill else _xml(f"<a:noFill {NSDECL}/>"))
                tcpr.set("anchor", "ctr")


INCH_EMU = 914400
