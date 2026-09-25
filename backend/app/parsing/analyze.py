"""Структурный анализ слайдов шаблона (детерминированный).

Для каждого слайда-примера определяет:
* заголовок, кикер, подложку заголовка;
* «хром» (фон, логотипы, колонтитулы, линейки, повторяющиеся брендовые элементы) и контент;
* повторяющиеся группы (карточки / строки / шаги) и роли их текстов;
* вид слайда (title / section / agenda / content / closing) и признак служебного слайда.
"""
from __future__ import annotations

import re
import statistics
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

from .extract import ShapeInfo, SlideExtractor
from .model import Box, ItemGroup, Slot, TextStyle
from .ooxml import EMU_PER_INCH, Theme, first, xp

INCH = EMU_PER_INCH

SERVICE_WORDS = re.compile(
    r"(шаблон|используй|используйте|hex\s*#|#[0-9a-f]{6}\b|шрифт:|цвета:|иконки|логотип|https?://|www\.|"
    r"padding|font-family|горячие клавиши|оформлени[ея]\b|рекомендуем|обязательный блок|template|lorem)",
    re.I,
)
PLACEHOLDER_WORDS = re.compile(r"(вставить|вставьте|insert|qr[- ]?код|фото\b|photo|image here|логотип)", re.I)
NUMERIC = re.compile(r"^\s*([№#]?\d{1,3}([.,]\d+)?\s*[%xх×]?|0\d)\s*$")
CLOSING_WORDS = re.compile(r"(спасибо|благодар|thank|вопросы|контакты|questions|итоги|что важно запомнить)", re.I)
AGENDA_WORDS = re.compile(r"(содержани|agenda|план презентации|оглавлени|структура|contents)", re.I)
SECTION_LAYOUT = re.compile(r"(раздел|section|divider|разделител|глав)", re.I)
TITLE_LAYOUT = re.compile(r"(титул|title slide|обложк|cover)", re.I)
CLOSING_LAYOUT = re.compile(r"(спасибо|thank|end|финал|конец)", re.I)
AGENDA_LAYOUT = re.compile(r"(содержани|agenda|оглавл)", re.I)


@dataclass
class TopItem:
    top: ShapeInfo
    desc: list[ShapeInfo]

    @property
    def sid(self) -> str:
        return self.top.sid

    @property
    def text(self) -> str:
        if self.top.kind != "group":
            return self.top.text
        return "\n".join(d.text for d in self.desc if d.text)

    @property
    def has_text(self) -> bool:
        return bool(self.text.strip())

    def texts(self) -> list[ShapeInfo]:
        if self.top.kind != "group":
            return [self.top] if self.top.has_text else []
        return [d for d in self.desc if d.has_text]

    @property
    def max_size(self) -> float:
        sizes = [p.style.size or 0 for t in self.texts() for p in t.paras]
        return max(sizes) if sizes else 0.0


@dataclass
class SlideAnalysis:
    index: int
    layout: str
    items: list[TopItem]
    shapes: list[ShapeInfo]
    title: Optional[TopItem] = None
    kicker: Optional[TopItem] = None
    label: Optional[TopItem] = None
    chrome: set[str] = field(default_factory=set)
    content: set[str] = field(default_factory=set)
    removable: set[str] = field(default_factory=set)       # плейсхолдеры фото и т.п.
    groups: list[ItemGroup] = field(default_factory=list)
    kind: str = "content"
    service: bool = False
    service_reason: str = ""
    band_top: int = 0
    band_bottom: int = 0
    features: dict = field(default_factory=dict)
    slots: list[Slot] = field(default_factory=list)


def _style_of(s: ShapeInfo) -> TextStyle:
    st = s.style
    ls = s.paras[0].line_spacing if s.paras else 1.0
    return TextStyle(font=st.font or "Arial", size=st.size or 14, color=st.color or "000000",
                     bold=bool(st.bold), caps=bool(st.caps), spacing=st.spacing, line_spacing=ls)


def _contains(outer: ShapeInfo, inner: ShapeInfo, tol: int = 0) -> bool:
    cx, cy = inner.x + inner.w / 2, inner.y + inner.h / 2
    return outer.x - tol <= cx <= outer.x2 + tol and outer.y - tol <= cy <= outer.y2 + tol


def _max_chars(s: ShapeInfo, size: float) -> int:
    w_pt = max(1.0, (s.w - s.insets[0] - s.insets[2]) / 12700)
    h_pt = max(1.0, (s.h - s.insets[1] - s.insets[3]) / 12700)
    per_line = w_pt / (size * 0.55)
    lines = max(1, int(h_pt / (size * 1.2)))
    return int(per_line * lines)


class TemplateAnalyzer:
    def __init__(self, path: Path):
        self.path = Path(path)
        self.prs = Presentation(str(path))
        self.W = self.prs.slide_width
        self.H = self.prs.slide_height
        self._themes: dict[int, Theme] = {}
        self.extractor = SlideExtractor(self.prs, self.theme_for_master)
        self.slides: list[SlideAnalysis] = []

    # --------------------------------------------------------------- themes
    def theme_for_master(self, master_el) -> Theme:
        key = id(master_el)
        if key in self._themes:
            return self._themes[key]
        theme_el = None
        clr_map = None
        for m in self.prs.slide_masters:
            if m._element is master_el:
                try:
                    tp = m.part.part_related_by(RT.THEME)
                    theme_el = etree.fromstring(tp.blob)
                except Exception:
                    theme_el = None
                cm = first(master_el, "p:clrMap")
                if cm is not None:
                    clr_map = dict(cm.attrib)
        th = Theme(theme_el, clr_map)
        self._themes[key] = th
        return th

    @property
    def main_theme(self) -> Theme:
        return self.theme_for_master(self.prs.slide_masters[0]._element)

    # --------------------------------------------------------------- run
    def run(self) -> list[SlideAnalysis]:
        for i, slide in enumerate(self.prs.slides, 1):
            layout = slide.slide_layout
            master = layout.slide_master
            shapes = self.extractor.extract(slide._element, slide.part, layout._element, master._element)
            items: list[TopItem] = []
            by_top: dict[int, list[ShapeInfo]] = defaultdict(list)
            for s in shapes:
                by_top[s.top_index].append(s)
            for ti in sorted(by_top):
                group = by_top[ti]
                top = next(s for s in group if s.depth == 0)
                items.append(TopItem(top=top, desc=[s for s in group if s.depth > 0]))
            self.slides.append(SlideAnalysis(index=i, layout=layout.name or "", items=items, shapes=shapes))
        recurring = self._recurring_signatures()
        for sa in self.slides:
            self._analyze_slide(sa, recurring)
        # второй проход: вид слайда с учётом медианного кегля заголовков колоды
        tsz = [sa.features["title_size"] for sa in self.slides if not sa.service and sa.features.get("title_size")]
        self.median_title = statistics.median(tsz) if tsz else 24
        for sa in self.slides:
            sa.kind = self._classify(sa)
            sa.slots = self._slots(sa)
        return self.slides

    def _sig(self, s: ShapeInfo) -> tuple:
        W, H = self.W, self.H
        t = s.text.strip() if len(s.text.strip()) <= 3 else "T" if s.has_text else ""
        return (s.kind, round(s.x / W, 2), round(s.y / H, 2), round(s.w / W, 2), round(s.h / H, 2), t)

    def _recurring_signatures(self) -> set:
        cnt: Counter = Counter()
        for sa in self.slides:
            seen = set()
            for it in sa.items:
                sig = self._sig(it.top)
                if sig not in seen:
                    cnt[sig] += 1
                    seen.add(sig)
        n = len(self.slides)
        return {sig for sig, c in cnt.items() if c >= max(3, int(0.3 * n)) and sig[0] != "table"}

    # --------------------------------------------------------------- slide
    def _analyze_slide(self, sa: SlideAnalysis, recurring: set) -> None:
        W, H = self.W, self.H
        slide_area = W * H
        items = sa.items

        # ---- служебный слайд
        all_text = " ".join(it.text for it in items)
        n_textless = sum(1 for s in sa.shapes if not s.has_text and s.kind not in ("group",))
        svc_hits = SERVICE_WORDS.findall(all_text)
        mono = any((s.style.font or "").lower() in ("consolas", "courier new", "jetbrains mono", "menlo")
                   for s in sa.shapes if s.has_text)
        if len(svc_hits) >= 2 or (svc_hits and len(all_text) < 400):
            sa.service, sa.service_reason = True, f"служебные слова: {', '.join(sorted(set(h.lower() for h in svc_hits))[:4])}"
        elif mono and ("{" in all_text or ";" in all_text):
            sa.service, sa.service_reason = True, "пример кода/служебный"
        elif n_textless > 120 and len(all_text) < 300:
            sa.service, sa.service_reason = True, "библиотека элементов (много фигур без текста)"

        # ---- заголовок
        sa.title = self._find_title(sa)
        title = sa.title
        if title is not None:
            t = title.top
            sa.band_top = t.y
            sa.band_bottom = t.y2
            # кикер над заголовком
            tsize = title.max_size or (t.style.size or 24)
            for it in items:
                if it is title or not it.has_text or it.top.kind == "group":
                    continue
                s = it.top
                if s.y2 <= t.y + 0.15 * t.h and s.y < 0.2 * H and (s.style.size or 12) < 0.75 * tsize and len(it.text) < 60:
                    sa.kicker = it
                    sa.band_top = min(sa.band_top, s.y)
                    break
            # подложка заголовка (плашка)
            for it in items:
                s = it.top
                if it is title or it.has_text or s.kind not in ("shape",) or s.fill in (None, "none", "picture"):
                    continue
                v_overlap = max(0, min(s.y2, t.y2) - max(s.y, t.y))
                if (v_overlap >= 0.6 * min(t.h, s.h) and s.x <= t.x + 0.25 * INCH and s.x2 > t.x + 0.3 * INCH
                        and s.h <= 2.8 * max(t.h, 1) and s.w < 0.92 * W and s.area < 0.2 * slide_area):
                    sa.label = it
                    sa.band_top = min(sa.band_top, s.y)
                    sa.band_bottom = max(sa.band_bottom, s.y2)
                    break

        # ---- хром
        chrome: set[str] = set()
        footer_y = 0.86 * H
        for it in items:
            s = it.top
            if title is not None and it is title:
                continue
            if sa.kicker is it or sa.label is it:
                continue
            is_bg = not it.has_text and s.area >= 0.6 * slide_area and s.kind in ("picture", "shape", "group")
            is_ph_chrome = s.ph_type in ("sldNum", "ftr", "dt")
            # колонтитул — мелкий текст; крупная подпись внизу (элемент схемы) — контент
            in_footer = s.y >= footer_y and s.h <= 0.12 * H and (not it.has_text or (len(it.text) < 40 and it.max_size <= 12.5))
            small_corner = (s.area < 0.035 * slide_area and not it.has_text and s.kind in ("picture", "group")
                            and (s.y2 < 0.14 * H or s.y > 0.86 * H))
            page_num = it.has_text and NUMERIC.match(it.text or "") and s.y > 0.85 * H and s.area < 0.01 * slide_area
            rule = (not it.has_text and s.kind in ("line", "shape") and min(s.h, s.w) < 0.012 * H
                    and s.w > 0.3 * W and title is not None
                    and (title.top.y2 - 0.02 * H <= s.y <= title.top.y2 + 0.14 * H))
            rec = self._sig(s) in recurring and (not it.has_text or len(it.text) < 4)
            if is_bg or is_ph_chrome or in_footer or small_corner or page_num or rule or rec:
                chrome.add(s.sid)
                if rule:
                    sa.band_bottom = max(sa.band_bottom, s.y2)
        sa.chrome = chrome

        # ---- контент
        keep = set(chrome)
        for it in (title, sa.kicker, sa.label):
            if it is not None:
                keep.add(it.sid)
        sa.content = {it.sid for it in items if it.sid not in keep}
        sa.removable = {it.sid for it in items if it.sid in sa.content and PLACEHOLDER_WORDS.search(it.text or "")
                        and len(it.text) < 40}

        # ---- признаки
        cont = [it for it in items if it.sid in sa.content]
        texts = [it for it in cont if it.has_text]
        f = {
            "n_content": len(cont),
            "n_text": len(texts),
            "chars": sum(len(it.text) for it in texts),
            "has_table": any(it.top.kind == "table" for it in cont),
            "has_chart": any(it.top.kind == "chart" for it in cont),
            "has_smartart": any(it.top.kind == "smartart" for it in cont),
            "big_picture": any(it.top.kind == "picture" and it.top.area > 0.12 * slide_area for it in cont),
            "big_numbers": sum(1 for it in texts if it.max_size >= 36 and NUMERIC.match(it.text.strip() or "x")),
            "bullets": sum(1 for it in texts for t in it.texts() for p in t.paras if p.bullet),
            "title_size": title.max_size if title else 0,
        }
        sa.features = f
        sa.groups = self._find_groups(sa)
        f["n_items"] = max((g.count for g in sa.groups), default=0)

    def _find_title(self, sa: SlideAnalysis) -> Optional[TopItem]:
        H = self.H
        phs = [it for it in sa.items if it.top.ph_type in ("title", "ctrTitle")]
        if phs:
            return phs[0]
        cands = []
        for it in sa.items:
            s = it.top
            if not it.has_text or s.kind == "group" or s.kind == "table":
                continue
            if NUMERIC.match(it.text) or s.y > 0.45 * H:
                continue
            size = it.max_size
            score = size
            if re.search(r"(title|заголов|heading)", s.name, re.I):
                score *= 1.35
            if s.y < 0.25 * H:
                score *= 1.2
            if len(it.text) > 160:
                score *= 0.5
            cands.append((score, -s.y, it))
        if not cands:
            return None
        cands.sort(key=lambda c: (c[0], c[1]), reverse=True)
        best = cands[0][2]
        # медиана по остальным текстам: единственный крупный текст (разделитель) не сравниваем сам с собой
        sizes = [p.style.size or 12 for s in sa.shapes if s.has_text and s.sid != best.top.sid for p in s.paras]
        med = statistics.median(sizes) if sizes else 12
        if best.max_size < max(14, med * 1.1) and best.top.y > 0.25 * H:
            return None
        return best

    # --------------------------------------------------------------- groups
    def _find_groups(self, sa: SlideAnalysis) -> list[ItemGroup]:
        W, H = self.W, self.H
        sarea = W * H
        cont = [it for it in sa.items if it.sid in sa.content]
        containers = [it for it in cont if not it.has_text and it.top.kind == "shape"
                      and (it.top.fill not in (None, "none") or it.top.line)
                      and 0.008 * sarea <= it.top.area <= 0.45 * sarea]
        groups: list[ItemGroup] = []
        members: dict[str, list[TopItem]] = defaultdict(list)
        assigned: set[str] = set()
        others = [it for it in cont if it not in containers] + [it for it in containers]
        for it in others:
            best = None
            for c in containers:
                if c is it or c.top.area <= it.top.area * 1.05:
                    continue
                if _contains(c.top, it.top):
                    if best is None or c.top.area < best.top.area:
                        best = c
            if best is not None:
                members[best.sid].append(it)
                assigned.add(it.sid)
        # кластеризация контейнеров по размеру
        with_members = [c for c in containers if members.get(c.sid) and c.sid not in assigned]
        clusters: list[list[TopItem]] = []
        for c in sorted(with_members, key=lambda c: (c.top.y, c.top.x)):
            for cl in clusters:
                r = cl[0].top
                if abs(r.w - c.top.w) <= 0.1 * r.w and abs(r.h - c.top.h) <= 0.1 * r.h:
                    cl.append(c)
                    break
            else:
                clusters.append([c])
        for cl in clusters:
            if len(cl) < 2:
                continue
            g = ItemGroup(container_sids=[c.sid for c in cl], count=len(cl))
            g.arrangement = self._arrangement([c.top for c in cl])
            for c in self._order(cl, g.arrangement):
                g.items.append({
                    "container": c.sid,
                    "box": _box(c.top).model_dump(),
                    "members": _assign_roles([self._member(m, c.top) for m in self._flatten(members[c.sid], members)]),
                })
            groups.append(g)
        if groups:
            groups.sort(key=lambda g: g.count * sum(len(i["members"]) for i in g.items), reverse=True)
            return groups
        # запасной путь: пункты без подложек — наборы текстов одного стиля, выровненные в ряд/колонку/сетку
        return self._text_groups(cont)

    @staticmethod
    def _bins(vals: list[float], tol: float) -> int:
        vals = sorted(vals)
        n = 1
        for a, b in zip(vals, vals[1:]):
            if b - a > tol:
                n += 1
        return n

    def _regular(self, v: list[TopItem]) -> Optional[str]:
        """Расположение набора, если оно регулярное (ряд, колонка, полная сетка без наложений), иначе None."""
        W, H = self.W, self.H
        tops = [it.top for it in v]
        rows = self._bins([s.y + s.h / 2 for s in tops], 0.08 * H)
        cols = self._bins([s.x + s.w / 2 for s in tops], 0.08 * W)
        n = len(v)
        for i, a in enumerate(tops):
            for b in tops[i + 1:]:
                ix = min(a.x2, b.x2) - max(a.x, b.x)
                iy = min(a.y2, b.y2) - max(a.y, b.y)
                if ix > 0.05 * min(a.w, b.w) and iy > 0.05 * min(a.h, b.h):
                    return None
        if rows == 1 and cols == n:
            return "row"
        if cols == 1 and rows == n:
            return "column"
        if rows >= 2 and cols >= 2 and rows * cols == n:
            return "grid"
        return None

    def _text_groups(self, cont: list[TopItem]) -> list[ItemGroup]:
        W, H = self.W, self.H
        texts = [it for it in cont if it.has_text and it.top.kind not in ("group", "table")]
        sets: dict[tuple, list[TopItem]] = defaultdict(list)
        for it in texts:
            sets[("t", round(it.max_size), bool(it.top.style.bold), bool(NUMERIC.match(it.text)))].append(it)
        # мелкие картинки близкого размера (иконки пунктов)
        pics = sorted((it for it in cont if it.top.kind == "picture" and it.top.area < 0.05 * W * H),
                      key=lambda it: it.top.area)
        for it in pics:
            for key, v in sets.items():
                r = v[0].top
                if key[0] == "p" and abs(r.w - it.top.w) <= 0.25 * r.w and abs(r.h - it.top.h) <= 0.25 * r.h:
                    v.append(it)
                    break
            else:
                sets[("p", len(sets))].append(it)
        cand = [(k, v, self._regular(v)) for k, v in sets.items() if k[0] == "t" and 2 <= len(v) <= 6]
        cand = [c for c in cand if c[2]]
        if not cand:
            return []
        # якорь — самый многочисленный регулярный набор; при равенстве — жирные ненумерованные (заголовки пунктов)
        key0, anchors, arr = max(cand, key=lambda c: (len(c[1]), c[0][2], not c[0][3], c[0][1]))
        n = len(anchors)
        anchors = self._order(anchors, arr)
        centers = [(a.top.x + a.top.w / 2, a.top.y + a.top.h / 2) for a in anchors]
        items: list[list[TopItem]] = [[a] for a in anchors]
        for key, v in sets.items():
            if key == key0 or len(v) != n:
                continue
            ordered = self._order(v, arr)
            ok = True
            for k, m in enumerate(ordered):
                cx, cy = m.top.x + m.top.w / 2, m.top.y + m.top.h / 2
                d = [((cx - ax) ** 2 + (cy - ay) ** 2) for ax, ay in centers]
                if min(range(n), key=lambda j: d[j]) != k:
                    ok = False
                    break
            if ok:
                for k, m in enumerate(ordered):
                    items[k].append(m)
        g = ItemGroup(count=n, arrangement=arr)
        for mem in items:
            xs = [m.top for m in mem]
            box = Box(x=min(s.x for s in xs), y=min(s.y for s in xs),
                      w=max(s.x2 for s in xs) - min(s.x for s in xs), h=max(s.y2 for s in xs) - min(s.y for s in xs))
            ms = [self._member(m, None) for m in mem]
            for m, src in zip(ms, mem):
                if src.top.kind == "picture":
                    m["role"] = "icon"
            g.items.append({"container": None, "box": box.model_dump(), "members": _assign_roles(ms)})
        return [g]

    def _flatten(self, mem: list[TopItem], members: dict) -> list[TopItem]:
        out = []
        for m in mem:
            out.append(m)
            if m.sid in members:
                out.extend(self._flatten(members[m.sid], members))
        return out

    def _member(self, m: TopItem, container: Optional[ShapeInfo]) -> dict:
        s = m.top
        role = "shape"
        empty_slot = (not m.has_text and s.text_slot and s.ph_type is not None
                      and s.w > 0.4 * 914400 and s.h > 0.15 * 914400)
        if empty_slot:
            return {"sid": m.sid, "role": "text", "box": _box(s).model_dump(), "text": "",
                    "style": _style_of(s).model_dump(), "insets": list(s.insets)}
        if m.has_text:
            t = m.text.strip()
            if NUMERIC.match(t):
                role = "number"
            elif s.style.bold or len(t) < 40:
                role = "heading"
            else:
                role = "body"
        elif s.kind == "picture":
            role = "icon" if container is not None and s.area < 0.15 * container.area else "image"
        elif container is not None and s.area < 0.08 * container.area:
            role = "marker"
        return {"sid": m.sid, "role": role, "box": _box(s).model_dump(), "text": m.text[:200],
                "style": _style_of(m.texts()[0]).model_dump() if m.has_text else None, "insets": list(s.insets)}

    def _arrangement(self, shapes: list[ShapeInfo]) -> str:
        ys = {round((s.y + s.h / 2) / (0.1 * self.H)) for s in shapes}
        xs = {round((s.x + s.w / 2) / (0.06 * self.W)) for s in shapes}
        if len(ys) == 1:
            return "row"
        if len(xs) == 1:
            return "column"
        return "grid"

    def _order(self, items: list[TopItem], arr: str) -> list[TopItem]:
        if arr == "column":
            return sorted(items, key=lambda i: i.top.y)
        if arr == "row":
            return sorted(items, key=lambda i: i.top.x)
        return sorted(items, key=lambda i: (round(i.top.y / (0.1 * self.H)), i.top.x))

    # --------------------------------------------------------------- kind
    def _classify(self, sa: SlideAnalysis) -> str:
        f = sa.features
        lay = sa.layout
        title_text = sa.title.text if sa.title else ""
        all_text = " ".join(it.text for it in sa.items)
        if sa.index == 1 or TITLE_LAYOUT.search(lay):
            if f["n_text"] <= 4 and not f["has_table"] and not f["has_chart"]:
                return "title"
        if CLOSING_LAYOUT.search(lay) or (CLOSING_WORDS.search(title_text) and f["n_text"] <= 6):
            return "closing"
        if AGENDA_WORDS.search(title_text) or (AGENDA_LAYOUT.search(lay) and sa.index <= 4):
            return "agenda"
        big_num = any(it.max_size >= 60 and NUMERIC.match(it.text.strip() or "x") for it in sa.items)
        med = getattr(self, "median_title", 24)
        if SECTION_LAYOUT.search(lay) or (f["n_text"] <= 4 and not f["has_table"] and not f["has_chart"]
                                          and not f.get("n_items") and sa.title is not None
                                          and (big_num or f["title_size"] >= 1.25 * med)):
            return "section"
        if sa.index == len(self.slides) and CLOSING_WORDS.search(all_text):
            return "closing"
        return "content"

    # --------------------------------------------------------------- slots
    def _slots(self, sa: SlideAnalysis) -> list[Slot]:
        slots: list[Slot] = []
        if sa.title is not None:
            t = sa.title.top
            st = _style_of(t)
            slots.append(Slot(role="title", sid=t.sid, box=_box(t), style=st,
                              max_chars=_max_chars(t, st.size), text=sa.title.text[:200]))
        if sa.kicker is not None:
            k = sa.kicker.top
            slots.append(Slot(role="kicker", sid=k.sid, box=_box(k), style=_style_of(k),
                              max_chars=_max_chars(k, k.style.size or 10), text=sa.kicker.text))
        for it in sa.items:
            s = it.top
            if it.sid in sa.chrome and it.has_text and NUMERIC.match(it.text) and s.y > 0.85 * self.H:
                slots.append(Slot(role="pagenum", sid=s.sid, box=_box(s), style=_style_of(s), text=it.text))
            elif s.ph_type == "sldNum":
                slots.append(Slot(role="pagenum", sid=s.sid, box=_box(s), style=_style_of(s), text=it.text))
        if sa.kind in ("title", "section", "closing"):
            rest = [it for it in sa.items if it.sid in sa.content and it.has_text and it.top.kind != "group"
                    and it.sid not in sa.removable]
            rest.sort(key=lambda it: (-it.max_size, it.top.y))
            used_sub = False
            has_kicker = any(sl.role == "kicker" for sl in slots)
            tsz = sa.title.max_size if sa.title else 24
            for it in rest:
                s = it.top
                st = _style_of(s)
                multi = len([ln for ln in it.text.splitlines() if ln.strip()]) >= 2
                above = sa.title is not None and s.y2 <= sa.title.top.y + 0.2 * sa.title.top.h
                if not has_kicker and above and (st.size or 12) < 0.6 * tsz and len(it.text) < 60 and not multi:
                    slots.append(Slot(role="kicker", sid=s.sid, box=_box(s), style=st,
                                      max_chars=_max_chars(s, st.size), text=it.text))
                    has_kicker = True
                    continue
                if NUMERIC.match(it.text) and it.max_size >= 40:
                    role = "number"
                elif multi and sa.kind in ("section", "closing"):
                    role = "body"
                elif not used_sub and (st.size or 0) >= 12 and len(it.text) < 200:
                    role = "subtitle"
                    used_sub = True
                elif len(it.text.split("\n")) >= 2 and sa.kind in ("section", "closing"):
                    role = "body"
                else:
                    role = "meta"
                slots.append(Slot(role=role, sid=s.sid, box=_box(s), style=st,
                                  max_chars=_max_chars(s, st.size), text=it.text[:300]))
        return slots


def _assign_roles(members: list[dict]) -> list[dict]:
    """Роли текстов внутри пункта: number / heading / body / extra (по порядку, кеглю и содержимому)."""
    texts = [m for m in members if m["role"] in ("heading", "body", "number", "text")]
    texts.sort(key=lambda m: (m["box"]["y"], m["box"]["x"]))
    have_head = False
    have_body = False
    for m in texts:
        t = (m.get("text") or "").strip()
        size = (m.get("style") or {}).get("size") or 12
        if t and NUMERIC.match(t):
            m["role"] = "number"
        elif not t and size >= 28 and m["box"]["w"] < 1.6 * 914400:
            m["role"] = "number"
        elif not have_head and (len(t) < 45 or (m.get("style") or {}).get("bold")) and m["box"]["h"] < 1.0 * 914400:
            m["role"], have_head = "heading", True
        elif not have_body:
            m["role"], have_body = "body", True
        else:
            m["role"] = "extra"
    if not have_body:
        # единственный текстовый слот без заголовка — это тело
        for m in texts:
            if m["role"] == "heading" and len([x for x in texts if x["role"] == "heading"]) == 1 and m["box"]["h"] > 0.5 * 914400:
                m["role"] = "body"
    return members


def _box(s: ShapeInfo) -> Box:
    return Box(x=int(s.x), y=int(s.y), w=int(s.w), h=int(s.h))
