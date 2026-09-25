"""Холсты: слайды шаблона без контента + пиксельный анализ свободной области и фона."""
from __future__ import annotations

import re
from pathlib import Path
from typing import Optional

import numpy as np
from PIL import Image
from pptx import Presentation

from ..core import pptx_ops
from .analyze import SlideAnalysis, TemplateAnalyzer
from .model import Box, Canvas
from .ooxml import luminance, rgb_to_hex
from .render import render_deck


def _layout_shapes(an: TemplateAnalyzer, cache: dict, slide_index: int):
    slide = an.prs.slides[slide_index - 1]
    lay = slide.slide_layout
    key = id(lay._element)
    if key not in cache:
        try:
            cache[key] = an.extractor.extract(lay._element, lay.part, None, lay.slide_master._element, in_layout=True)
        except Exception:
            cache[key] = []
    return cache[key]


def _intersect(a: Box, x: int, y: int, w: int, h: int) -> int:
    ix = max(0, min(a.x2, x + w) - max(a.x, x))
    iy = max(0, min(a.y2, y + h) - max(a.y, y))
    return ix * iy


def _removable_containers(sa: SlideAnalysis) -> set[str]:
    """Подложки плейсхолдеров «Вставить фото/QR» удаляются вместе с подписью."""
    out = set()
    for rid in sa.removable:
        r = next((it.top for it in sa.items if it.sid == rid), None)
        if r is None:
            continue
        for it in sa.items:
            s = it.top
            if it.sid == rid or it.has_text:
                continue
            cx, cy = r.x + r.w / 2, r.y + r.h / 2
            if s.x <= cx <= s.x2 and s.y <= cy <= s.y2 and s.area <= 6 * max(1, r.area):
                out.add(it.sid)
    return out


def _master_shapes(an: TemplateAnalyzer, cache: dict, slide_index: int):
    slide = an.prs.slides[slide_index - 1]
    m = slide.slide_layout.slide_master
    key = ("m", id(m._element))
    if key not in cache:
        try:
            cache[key] = an.extractor.extract(m._element, m.part, None, m._element, in_layout=True)
        except Exception:
            cache[key] = []
    return cache[key]


def _title_max_w(an: TemplateAnalyzer, sa: SlideAnalysis, lay_cache: dict, gap: int) -> Optional[int]:
    """Правая граница заголовка: логотипы и прочий хром на одной высоте с заголовком."""
    if sa.title is None:
        return None
    t = sa.title.top
    obstacles = []
    for it in sa.items:
        if it.sid in sa.chrome:
            obstacles.append(it.top)
    shapes = _layout_shapes(an, lay_cache, sa.index)
    # макет может скрывать графику мастера (showMasterSp="0")
    lay_el = an.prs.slides[sa.index - 1].slide_layout._element
    if lay_el.get("showMasterSp") != "0":
        shapes = shapes + _master_shapes(an, lay_cache, sa.index)
    obstacles += [s for s in shapes if s.ph_type is None and s.depth == 0 and s.area < 0.5 * an.W * an.H]
    lim = None
    for o in obstacles:
        v = min(o.y2, t.y2) - max(o.y, t.y)
        if v > -int(0.08 * 914400) and o.x > t.x + int(0.5 * 914400) and o.x < t.x2 and o.h < 0.3 * an.H:
            cand = o.x - t.x - int(0.5 * gap)
            lim = cand if lim is None else min(lim, cand)
    return lim


TOKEN_TEXT = re.compile(r"\{\{[^{}]+\}\}")


def draft_canvases(an: TemplateAnalyzer, slides: list[SlideAnalysis], grid: Box, gap: int) -> list[Canvas]:
    W, H = an.W, an.H
    res: list[Canvas] = []
    seen: dict[tuple, Canvas] = {}
    lay_cache: dict = {}
    for sa in slides:
        if sa.service:
            continue
        if sa.kind in ("title", "section", "closing"):
            if not any(sl.role == "title" for sl in sa.slots):
                continue
            drop = sa.removable | _removable_containers(sa)
            # «стена логотипов» (≥3 небольших картинки) — иллюстрация темы образца, а не оформление
            pics = [it for it in sa.items if it.sid in sa.content and it.top.kind == "picture"
                    and it.top.area < 0.06 * W * H]
            if len(pics) >= 3:
                drop |= {it.sid for it in pics}
            # метки-заготовки {{...}} вне слотов — служебный текст шаблона, в колоду он не должен попасть
            slot_sids = {sl.sid for sl in sa.slots}
            drop |= {it.sid for it in sa.items if it.has_text and it.sid not in slot_sids and TOKEN_TEXT.search(it.text)}
            # пустые места под фото (серые круги у подписи спикера): без фотографии они выглядят как дыра
            from .analyze import is_photo_placeholder
            drop |= {it.sid for it in sa.items if it.sid in sa.content and is_photo_placeholder(it.top, W * H)}
            keep = [it.sid for it in sa.items if it.sid not in drop]
        else:
            if sa.title is None:
                continue
            keep = [it.sid for it in sa.items if it.sid in sa.chrome or it in (sa.title, sa.kicker, sa.label)]
        # контентная область
        footer_tops = [it.top.y for it in sa.items if it.sid in sa.chrome and it.top.y > 0.78 * H
                       and it.top.area < 0.3 * W * H]
        top = sa.band_bottom + max(gap, int(0.04 * H)) if sa.title is not None else grid.y
        bottom = min([grid.y2 + int(0.02 * H), H - int(0.06 * H)] + [y - int(0.6 * gap) for y in footer_tops])
        left, right = grid.x, grid.x2
        if sa.kind != "content":
            top = max(top, int(0.12 * H))
        cbox = Box(x=left, y=int(top), w=max(1, right - left), h=max(1, int(bottom - top)))
        slots = [sl for sl in sa.slots]
        sig = (sa.kind, sa.layout, len(keep), round(cbox.y / (0.03 * H)),
               tuple(sorted(round(sl.box.x / (0.02 * W)) for sl in slots if sl.role == "title")))
        if sig in seen and sa.kind == "content":
            seen[sig].score += 1
            continue
        c = Canvas(id=f"s{sa.index}", source_slide=sa.index, kind=sa.kind, layout=sa.layout, keep_sids=keep,
                   slots=slots, title_label_sid=sa.label.sid if sa.label else None, content_box=cbox, score=1.0,
                   title_max_w=_title_max_w(an, sa, lay_cache, gap))
        # статичный текст/декор макета в контентной области
        if sa.kind in ("content", "agenda"):
            for ls in _layout_shapes(an, lay_cache, sa.index):
                if ls.ph_type is not None or ls.depth > 0:
                    continue
                inter = _intersect(cbox, ls.x, ls.y, ls.w, ls.h)
                if ls.has_text and inter > 0.3 * max(1, ls.area) and not ls.text.strip().isdigit():
                    c.clean, c.dirty_reason = False, f"текст макета: {ls.text.strip()[:30]}"
                    break
                if (not ls.has_text and ls.kind in ("picture", "shape", "group") and ls.area < 0.6 * W * H
                        and inter > 0.08 * cbox.area):
                    c.dirty_reason = "декор макета в контентной области"
        seen[sig] = c
        res.append(c)
    return res


def build_canvas_deck(template: Path, canvases: list[Canvas], out: Path) -> Path:
    prs = Presentation(str(template))
    src = list(prs.slides)
    made = []
    for c in canvases:
        s = pptx_ops.clone_slide(prs, src[c.source_slide - 1])
        keep = set(c.keep_sids)
        pptx_ops.remove_shapes(s, {pptx_ops.shape_id(el) for el in pptx_ops.top_shapes(s)} - keep)
        made.append(s)
    pptx_ops.delete_slides(prs, made)
    prs.save(str(out))
    return out


def _largest_free_rect(free: np.ndarray) -> tuple[int, int, int, int]:
    """Крупнейший прямоугольник из True в булевой матрице (гистограммный метод). (r0, c0, r1, c1)."""
    rows, cols = free.shape
    heights = np.zeros(cols, dtype=int)
    best = (0, 0, 0, 0, 0)
    for r in range(rows):
        heights = np.where(free[r], heights + 1, 0)
        stack: list[int] = []
        for c in range(cols + 1):
            h = heights[c] if c < cols else 0
            start = c
            while stack and heights[stack[-1]] >= h:
                top = stack.pop()
                height = heights[top]
                left = stack[-1] + 1 if stack else 0
                area = height * (c - left)
                if area > best[0]:
                    best = (area, r - height + 1, left, r + 1, c)
                start = top
            stack.append(c)
    return best[1], best[2], best[3], best[4]


def _pixel_title_limit(c: Canvas, sl, arr: np.ndarray, sx: float, sy: float, W: int, gap: int) -> None:
    """Графика, «вшитая» в фон (логотипы в картинке макета), справа от заголовка ограничивает его ширину."""
    from .fonts import text_width_pt
    ih, iw = arr.shape[:2]
    text_w = int(text_width_pt(sl.text.upper() if sl.style.caps else sl.text, sl.style.font, sl.style.size,
                               sl.style.bold) * 12700) if sl.text.strip() else 0
    y0 = max(0, int((sl.box.y - 0.25 * 914400) * sy))
    y1 = min(ih, int(sl.box.y2 * sy))
    # где на рендере начинается сам текст (заголовок бывает по центру или справа)
    text_x = sl.box.x + 91440
    bx0, bx1 = max(0, int(sl.box.x * sx)), min(iw, int(sl.box.x2 * sx))
    if text_w and bx1 - bx0 > 10 and y1 > y0:
        g = arr[max(0, int(sl.box.y * sy)):y1, bx0:bx1].mean(axis=2)
        gy_, gx_ = np.gradient(g)
        busy_cols = np.where((np.hypot(gx_, gy_) > 14).mean(axis=0) > 0.04)[0]
        if len(busy_cols):
            text_x = max(text_x, int((bx0 + busy_cols[0]) / sx))
    start = text_x + text_w + int(0.35 * 914400)
    x0, x1 = int(start * sx), min(iw, int(sl.box.x2 * sx))
    if x1 - x0 < 10 or y1 - y0 < 4:
        return
    gray = arr[y0:y1, x0:x1].mean(axis=2)
    gy, gx = np.gradient(gray)
    busy = np.hypot(gx, gy) > 14
    cols = busy.mean(axis=0) > 0.04
    run = 0
    need = max(3, int(0.04 * 914400 * sx))
    for i, b in enumerate(cols):
        run = run + 1 if b else 0
        if run >= need:
            ox = (x0 + i - run + 1) / sx
            lim = int(ox - sl.box.x - 0.5 * gap)
            if lim > 0.25 * sl.box.w:
                c.title_max_w = lim if c.title_max_w is None else min(c.title_max_w, lim)
            return


def analyze_canvas_images(canvases: list[Canvas], pngs: list[Path], W: int, H: int, gap: int) -> None:
    for c, png in zip(canvases, pngs):
        c.preview = str(png)
        img = Image.open(png).convert("RGB")
        iw, ih = img.size
        arr = np.asarray(img).astype(np.float32)
        sx, sy = iw / W, ih / H
        b = c.content_box
        x0, y0 = max(0, int(b.x * sx)), max(0, int(b.y * sy))
        x1, y1 = min(iw, int(b.x2 * sx)), min(ih, int(b.y2 * sy))
        region = arr[y0:y1, x0:x1] if (x1 > x0 and y1 > y0) else arr
        med = np.median(region.reshape(-1, 3), axis=0)
        c.bg = rgb_to_hex(med)
        for sl in c.slots:
            bx0, by0 = max(0, int(sl.box.x * sx)), max(0, int(sl.box.y * sy))
            bx1, by1 = min(iw, int(sl.box.x2 * sx)), min(ih, int(sl.box.y2 * sy))
            if bx1 > bx0 and by1 > by0:
                sl.bg = rgb_to_hex(np.median(arr[by0:by1, bx0:bx1].reshape(-1, 3), axis=0))
            if sl.role == "title" and c.kind in ("content", "agenda"):
                _pixel_title_limit(c, sl, arr, sx, sy, W, gap)
        c.dark = luminance(c.bg) < 0.25
        # «занятость» фона: градиенты допустимы, резкие края — нет
        gray = region.mean(axis=2)
        gy, gx = np.gradient(gray)
        mag = np.hypot(gx, gy)
        busy = mag > 10
        c.bg_busy = bool(busy.mean() > 0.02 or region.std(axis=(0, 1)).mean() > 28)
        edge_frac = float(busy.mean())
        if c.kind in ("content", "agenda") and edge_frac > 0.003:
            c.clean = False
            c.dirty_reason = c.dirty_reason or "декор в контентной области"
            cell = max(4, int(iw / 120))
            gh, gw = busy.shape[0] // cell, busy.shape[1] // cell
            if gh > 3 and gw > 3:
                blocks = busy[: gh * cell, : gw * cell].reshape(gh, cell, gw, cell).mean(axis=(1, 3)) > 0.02
                free = ~blocks
                r0, c0, r1, c1 = _largest_free_rect(free)
                area = (r1 - r0) * (c1 - c0)
                if area >= 0.45 * gh * gw and area < 0.97 * gh * gw:
                    nx = b.x + int(c0 * cell / sx)
                    ny = b.y + int(r0 * cell / sy)
                    nw = int((c1 - c0) * cell / sx)
                    nh = int((r1 - r0) * cell / sy)
                    pad = int(0.4 * gap)
                    c.content_box = Box(x=nx + (pad if c0 > 0 else 0), y=ny + (pad if r0 > 0 else 0),
                                        w=max(1, nw - (pad if c0 > 0 else 0) - (pad if c1 < gw else 0)),
                                        h=max(1, nh - (pad if r0 > 0 else 0) - (pad if r1 < gh else 0)))
                    c.bg_busy = False
                    # свободная часть достаточно велика — холст пригоден в «урезанном» виде
                    c.clean = area >= 0.6 * gh * gw and c.dirty_reason != "декор макета в контентной области" or area >= 0.8 * gh * gw
        # оценка пригодности холста
        ratio = c.content_box.area / max(1, W * H)
        c.score = c.score * (0.5 + ratio) * (0.6 if c.bg_busy else 1.0) * (1.0 if c.clean else 0.15)


def render_canvases(template: Path, canvases: list[Canvas], work: Path, W: int, H: int, gap: int) -> None:
    work.mkdir(parents=True, exist_ok=True)
    deck = build_canvas_deck(template, canvases, work / "canvases.pptx")
    _, pngs = render_deck(deck, work, prefix="canvas")
    analyze_canvas_images(canvases, pngs, W, H, gap)
