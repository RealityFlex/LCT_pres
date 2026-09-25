"""Растровые шаблоны: каждый слайд — одна картинка (экспорт из Figma/Canva), фигур и текста нет.

Из пикселей восстанавливается дизайн-система и собирается редактируемый шаблон-заменитель:
чистый фон (гладкая аппроксимация без контента), заголовки на своих местах, своим цветом и кеглем,
образец карточек и акцентов. Дальше он разбирается обычным пайплайном.
"""
from __future__ import annotations

import colorsys
import io
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

from . import fonts
from .analyze import SlideAnalysis, TemplateAnalyzer
from .ooxml import luminance, rgb_to_hex

WORK_W = 480   # ширина рабочего растра для анализа


def is_raster(an: TemplateAnalyzer, slides: list[SlideAnalysis]) -> bool:
    """Большинство слайдов без единого текста: одна картинка на весь слайд (или только фон)."""
    W, H = an.W, an.H
    n = 0
    for sa in slides:
        if any(s.has_text for s in sa.shapes):
            continue
        tops = [s for s in sa.shapes if s.depth == 0]
        if not tops or any(s.kind == "picture" and s.area >= 0.85 * W * H for s in tops):
            n += 1
    return bool(slides) and n >= max(1, 0.6 * len(slides))


# ------------------------------------------------------------------ пиксельный анализ

def _load(png: str) -> np.ndarray:
    im = Image.open(png).convert("RGB")
    h = int(WORK_W * im.height / im.width)
    return np.asarray(im.resize((WORK_W, h), Image.LANCZOS)).astype(np.float32)


def _design(xs: np.ndarray, ys: np.ndarray) -> np.ndarray:
    return np.stack([np.ones_like(xs), xs, ys, xs * xs, xs * ys, ys * ys], axis=1)


def fit_background(img: np.ndarray) -> np.ndarray:
    """Коэффициенты гладкой (квадратичной) поверхности фона по каналам; контент отсекается как выбросы."""
    h, w, _ = img.shape
    sy, sx = max(1, h // 54), max(1, w // 96)
    small = img[::sy, ::sx]
    gy, gx = np.mgrid[0:small.shape[0], 0:small.shape[1]]
    xs, ys = (gx.ravel() / small.shape[1]), (gy.ravel() / small.shape[0])
    A = _design(xs, ys)
    Y = small.reshape(-1, 3)
    # старт — рамка по краю слайда (там почти всегда чистый фон), затем расширяемся на похожие пиксели;
    # иначе крупные белые карточки «перетягивают» фон на себя
    ring = (xs < 0.05) | (xs > 0.95) | (ys < 0.04) | (ys > 0.96)
    keep = ring.copy()
    coef = np.zeros((6, 3))
    for _ in range(5):
        coef, *_ = np.linalg.lstsq(A[keep], Y[keep], rcond=None)
        res = np.linalg.norm(Y - A @ coef, axis=1)
        mad = np.median(res[ring & keep]) + 1e-3 if (ring & keep).any() else np.median(res) + 1e-3
        keep = res < max(8.0, 3.0 * mad)
        if keep.sum() < 0.05 * len(Y):
            keep = ring
    return coef


def bg_image(coef: np.ndarray, w: int, h: int) -> np.ndarray:
    gy, gx = np.mgrid[0:h, 0:w]
    A = _design(gx.ravel() / w, gy.ravel() / h)
    out = (A @ coef).reshape(h, w, 3)
    return np.clip(out, 0, 255)


def _bands(mask: np.ndarray, y0: int, y1: int, min_cov: float = 0.004) -> list[tuple[int, int, int, int]]:
    """Горизонтальные полосы текста: (y0, y1, x0, x1) в пикселях рабочего растра."""
    h, w = mask.shape
    rows = mask[y0:y1].sum(1) > max(2, min_cov * w)
    out = []
    y = 0
    while y < len(rows):
        if rows[y]:
            s = y
            while y < len(rows) and rows[y]:
                y += 1
            band = mask[y0 + s:y0 + y]
            cols = np.where(band.sum(0) > 0)[0]
            if len(cols):
                out.append((y0 + s, y0 + y, int(cols.min()), int(cols.max()) + 1))
        y += 1
    return out


@dataclass
class TextSpec:
    x: float          # доли слайда
    y: float
    w: float
    h: float
    size_frac: float  # высота строки глифов в долях высоты слайда
    color: str
    align: str        # l | ctr


@dataclass
class RasterDesign:
    bg_title: np.ndarray
    bg_content: np.ndarray
    dark: bool
    title_slide: Optional[TextSpec] = None
    subtitle_slide: Optional[TextSpec] = None
    title: Optional[TextSpec] = None
    subtitle: Optional[TextSpec] = None
    body_frac: float = 0.02
    body_color: str = "FFFFFF"
    card_fill: Optional[str] = None
    card_text: str = "1F1F1F"
    card_radius: float = 0.08
    accents: list[str] = field(default_factory=list)


def _fg_mask(img: np.ndarray, coef: np.ndarray) -> np.ndarray:
    h, w, _ = img.shape
    diff = np.linalg.norm(img - bg_image(coef, w, h), axis=2)
    return diff > 38


def _text_mask(mask: np.ndarray) -> np.ndarray:
    """Маска без крупных объектов (картинки, мокапы, карточки) — остаются строки текста на фоне."""
    from scipy import ndimage
    h, w = mask.shape
    lab, n = ndimage.label(ndimage.binary_dilation(mask, iterations=1))
    out = mask.copy()
    for i, sl in enumerate(ndimage.find_objects(lab), 1):
        if sl is None:
            continue
        if sl[0].stop - sl[0].start > 0.16 * h:
            out[sl][lab[sl] == i] = False
    return out


def _text_color(img: np.ndarray, band, mask: np.ndarray) -> str:
    y0, y1, x0, x1 = band
    px = img[y0:y1, x0:x1][mask[y0:y1, x0:x1]]
    if not len(px):
        return "FFFFFF"
    lum = px.mean(1)
    # глифы — самые контрастные к фону пиксели полосы (сглаженные края отбрасываем)
    bgl = np.median(img[y0:y1, x0:x1].reshape(-1, 3).mean(1))
    sel = px[lum >= np.percentile(lum, 70)] if lum.mean() > bgl else px[lum <= np.percentile(lum, 30)]
    return rgb_to_hex(np.median(sel, axis=0))


def _spec(img, mask, band) -> TextSpec:
    h, w, _ = img.shape
    y0, y1, x0, x1 = band
    cx = (x0 + x1) / 2 / w
    align = "ctr" if abs(cx - 0.5) < 0.06 and x0 > 0.12 * w else "l"
    return TextSpec(x=x0 / w, y=y0 / h, w=(x1 - x0) / w, h=(y1 - y0) / h, size_frac=(y1 - y0) / h,
                    color=_text_color(img, band, mask), align=align)


def analyze(pngs: list[str]) -> RasterDesign:
    imgs = [_load(p) for p in pngs]
    coefs = [fit_background(im) for im in imgs]
    content_coef = np.median(np.stack(coefs[1:] or coefs), axis=0)
    h, w, _ = imgs[0].shape
    mean_bg = bg_image(content_coef, 32, 18).reshape(-1, 3).mean(0)
    d = RasterDesign(bg_title=bg_image(coefs[0], 1920, 1080), bg_content=bg_image(content_coef, 1920, 1080),
                     dark=luminance(rgb_to_hex(mean_bg)) < 0.35)
    # ---- титул: самая высокая строка слайда 1 и строка под ней
    m0 = _text_mask(_fg_mask(imgs[0], coefs[0]))
    # строка текста вытянута по горизонтали (логотип — почти квадратный)
    b0 = [b for b in _bands(m0, 0, h) if 0.012 * h <= b[1] - b[0] <= 0.22 * h and (b[3] - b[2]) >= 3.0 * (b[1] - b[0])]
    if b0:
        big = max(b0, key=lambda b: b[1] - b[0])
        d.title_slide = _spec(imgs[0], m0, big)
        below = [b for b in b0 if b[0] >= big[1] and b[0] - big[1] < 0.12 * h and b[1] - b[0] < (big[1] - big[0]) * 0.8]
        if below:
            d.subtitle_slide = _spec(imgs[0], m0, below[0])
    # ---- контентные: первая строка в верхней полосе — заголовок, следующая мелкая — подзаголовок
    titles, subs, bodies = [], [], []
    body_cols: list[str] = []
    for im, cf in zip(imgs[1:], coefs[1:]):
        m = _text_mask(_fg_mask(im, cf))
        bands = [b for b in _bands(m, 0, int(0.25 * h)) if b[1] - b[0] >= 0.008 * h]
        if not bands or bands[0][1] - bands[0][0] < 0.018 * h:
            continue
        t = bands[0]
        titles.append(_spec(im, m, t))
        nxt = [b for b in bands[1:] if b[0] - t[1] < 0.06 * h and b[1] - b[0] < 0.7 * (t[1] - t[0])]
        if nxt:
            subs.append(_spec(im, m, nxt[0]))
        # текст основной области: низкие полосы в свободной (не занятой картинками) части
        for b in _bands(m, int(0.25 * h), h, 0.002):
            bh = b[1] - b[0]
            if 0.012 * h <= bh <= 0.035 * h and (b[3] - b[2]) > 0.15 * w:
                bodies.append(bh / h)
                body_cols.append(_text_color(im, b, m))
    if titles:
        k = int(np.argsort([t.size_frac for t in titles])[len(titles) // 2])
        d.title = titles[k]
    if subs:
        d.subtitle = subs[len(subs) // 2]
    if bodies:
        d.body_frac = float(np.median(bodies))
        d.body_color = max(set(body_cols), key=body_cols.count)
    # ---- карточки: крупные почти белые прямоугольные области
    from scipy import ndimage
    fills, radii = [], []
    for im in imgs[1:]:
        white = im.min(2) > 236
        lab, n = ndimage.label(white)
        for sl in ndimage.find_objects(lab):
            if sl is None:
                continue
            hh, ww = sl[0].stop - sl[0].start, sl[1].stop - sl[1].start
            if hh * ww < 0.012 * h * w or hh < 0.08 * h or ww < 0.08 * w:
                continue
            region = white[sl]
            if region.mean() < 0.8:
                continue
            fills.append(rgb_to_hex(np.median(im[sl][region], axis=0)))
            # скругление: сколько пикселей по диагонали угла не белые
            r = 0
            while r < min(hh, ww) // 3 and not region[r, r]:
                r += 1
            radii.append(r * 2.4 / min(hh, ww))
    if len(fills) >= 2:
        d.card_fill = max(set(fills), key=fills.count)
        d.card_radius = float(min(0.35, max(0.03, np.median(radii))))
    # ---- акценты: насыщенные цвета контента, далёкие от фона
    px = []
    for im, cf in zip(imgs, coefs):
        m = _fg_mask(im, cf)
        px.append(im[m][::7])
    allpx = np.concatenate(px) / 255.0 if px else np.zeros((0, 3))
    bins: dict[int, list] = {}
    bg_h = colorsys.rgb_to_hsv(*(mean_bg / 255.0))[0]
    for r, g, b in allpx[:60000]:
        hh, s, v = colorsys.rgb_to_hsv(r, g, b)
        if s < 0.45 or v < 0.45:
            continue
        bins.setdefault(int(hh * 18) % 18, []).append((r, g, b))
    ranked = sorted(bins.items(), key=lambda kv: -len(kv[1]))
    for hb, cols in ranked:
        if len(cols) < 40:
            break
        c = np.median(np.array(cols), axis=0) * 255
        hx = rgb_to_hex(c)
        hue = colorsys.rgb_to_hsv(*(c / 255))[0]
        if min(abs(hue - bg_h), 1 - abs(hue - bg_h)) < 0.04 and d.accents:
            continue
        d.accents.append(hx)
        if len(d.accents) >= 3:
            break
    return d


# ------------------------------------------------------------------ шаблон-заменитель

def _pick_font() -> str:
    for f in ("Manrope", "Montserrat", "Inter", "Arial"):
        if fonts.font_available(f) or (f != "Arial" and fonts.try_fetch_google_font(f)):
            return f
    return "Arial"


def _png(arr: np.ndarray) -> io.BytesIO:
    rng = np.random.default_rng(7)
    arr = np.clip(arr + rng.normal(0, 0.6, arr.shape), 0, 255)   # лёгкий шум против полос градиента
    buf = io.BytesIO()
    Image.fromarray(arr.astype(np.uint8)).save(buf, format="PNG")
    buf.seek(0)
    return buf


def _text(slide, x, y, w, h, text, font, size_pt, color, bold=False, align="l", anchor="t"):
    tb = slide.shapes.add_textbox(Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    tf.vertical_anchor = MSO_ANCHOR.BOTTOM if anchor == "b" else MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER if align == "ctr" else PP_ALIGN.LEFT
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(round(size_pt * 2) / 2)
    r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    return tb


def build_editable(src: Path, out: Path, d: RasterDesign) -> Path:
    prs0 = Presentation(str(src))
    W, H = prs0.slide_width, prs0.slide_height
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]

    def add(layout_idx: int = 6):
        # макеты «Title Slide» / «Section Header» подсказывают анализатору вид слайда; их плейсхолдеры убираем
        sl = prs.slides.add_slide(prs.slide_layouts[layout_idx])
        for ph in list(sl.placeholders):
            ph._element.getparent().remove(ph._element)
        return sl
    font = _pick_font()
    Hpt = H / 12700
    fg = "FFFFFF" if d.dark else "1F1F1F"
    ts = d.title or TextSpec(0.06, 0.06, 0.88, 0.07, 0.05, fg, "l")
    title_pt = max(20.0, min(40.0, ts.size_frac * Hpt / 0.85))
    body_pt = max(11.0, min(18.0, d.body_frac * Hpt / 0.85))
    sub_pt = max(10.0, min(title_pt * 0.6, (d.subtitle.size_frac * Hpt / 0.85) if d.subtitle else body_pt * 0.85))
    sub_col = d.subtitle.color if d.subtitle else d.body_color
    margin = int(0.06 * W)

    def bg(slide, arr):
        slide.shapes.add_picture(_png(arr), 0, 0, W, H)

    def title_box(slide, spec: TextSpec, text: str, size: float, sub: Optional[str] = None):
        line = size * 1.25 * 12700
        if spec.align == "ctr":
            x, w = margin, W - 2 * margin
        else:
            x, w = int(spec.x * W), int(W * 0.94) - int(spec.x * W)
        y = int(spec.y * H - 0.18 * line)
        _text(slide, x, y, w, int(line * 1.02), text, font, size, spec.color, bold=True, align=spec.align)
        if sub:
            _text(slide, x, y + int(line * 1.05), w, int(sub_pt * 1.4 * 12700), sub, font, sub_pt, sub_col,
                  align=spec.align)
        return y + line

    # 1. титул
    s = add(0)
    bg(s, d.bg_title)
    tsl = d.title_slide or TextSpec(0.06, 0.4, 0.5, 0.1, 0.1, fg, "l")
    big = max(title_pt * 1.4, min(60.0, tsl.size_frac * Hpt / 0.85))
    tw = TextSpec(tsl.x, tsl.y, max(tsl.w, 0.45), tsl.h, tsl.size_frac, tsl.color, tsl.align)
    y_end = title_box(s, tw, "Название презентации", big)
    sub_spec = d.subtitle_slide
    _text(s, int(tw.x * W) if tw.align == "l" else margin, int(y_end + 0.02 * H),
          int(W * 0.5) if tw.align == "l" else W - 2 * margin, int(0.12 * H), "Подзаголовок презентации", font,
          max(12.0, min(24.0, (sub_spec.size_frac * Hpt / 0.85) if sub_spec else big * 0.4)),
          sub_spec.color if sub_spec else sub_col, align=tw.align)
    # 2. разделитель
    s = add(2)
    bg(s, d.bg_content)
    sec = TextSpec(0.08, 0.38, 0.6, 0.1, 0.1, ts.color, "l")
    y_end = title_box(s, sec, "Раздел презентации", min(54.0, title_pt * 1.6))
    _text(s, int(0.08 * W), int(y_end + 0.03 * H), int(0.6 * W), int(0.1 * H), "Короткое пояснение к разделу",
          font, body_pt * 1.15, sub_col)
    # 3. контентный холст: без карточек в шаблоне — пустой слайд с заголовком;
    # иначе холстом служит слайд-образец карточек (его контент задаёт сетку шаблона)
    if not d.card_fill:
        s = add()
        bg(s, d.bg_content)
        title_box(s, ts, "Заголовок слайда", title_pt, "Пояснение к слайду" if d.subtitle else None)
    # 4. образец карточек
    if d.card_fill:
        s = add()
        bg(s, d.bg_content)
        y_end = title_box(s, ts, "Карточки", title_pt, "Пояснение к слайду" if d.subtitle else None)
        top = int(y_end + (0.1 if d.subtitle else 0.06) * H)
        gap = int(0.025 * W)
        cw = (W - 2 * margin - 2 * gap) // 3
        ch = int(0.42 * H)   # пропорции карточек исходных слайдов
        card_txt = "1F1F1F" if luminance(d.card_fill) > 0.5 else "FFFFFF"
        head_col = next((a for a in d.accents if _contrast_ok(a, d.card_fill)), card_txt)
        for i in range(3):
            x = margin + i * (cw + gap)
            card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x), Emu(top), Emu(cw), Emu(ch))
            card.adjustments[0] = d.card_radius
            card.fill.solid()
            card.fill.fore_color.rgb = RGBColor.from_string(d.card_fill)
            card.line.fill.background()
            card.shadow.inherit = False
            pad = int(0.08 * cw)
            yy = top + pad
            if d.accents:
                # маркер пункта в акцентном цвете шаблона (иконки/значки на исходных слайдах)
                mk = int(body_pt * 2.2 * 12700)
                dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Emu(x + pad), Emu(yy), Emu(mk), Emu(mk))
                dot.fill.solid()
                dot.fill.fore_color.rgb = RGBColor.from_string(d.accents[i % len(d.accents)])
                dot.line.fill.background()
                yy += mk + int(body_pt * 0.8 * 12700)
            _text(s, x + pad, yy, cw - 2 * pad, int(body_pt * 1.6 * 12700), f"Пункт {i + 1}", font,
                  body_pt * 1.2, head_col, bold=True)
            yy += int(body_pt * 2.0 * 12700)
            _text(s, x + pad, yy, cw - 2 * pad, top + ch - pad - yy,
                  "Короткое пояснение к пункту в одну-две строки", font, body_pt, card_txt)
    # 5. финал
    s = add()
    bg(s, d.bg_content)
    fin = TextSpec(0.1, 0.4, 0.8, 0.1, 0.1, ts.color, "ctr")
    title_box(s, fin, "Спасибо за внимание", min(54.0, title_pt * 1.6), "Контакты и вопросы")
    prs.save(str(out))
    return out


def _contrast_ok(fg: str, bg: str) -> bool:
    from .ooxml import contrast
    return contrast(fg, bg) >= 4.5
