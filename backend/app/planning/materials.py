"""Материалы к брифу: текст из приложенных файлов (README, документация, отчёты, старые презентации).

Извлечённый текст добавляется в «дополнительные данные» брифа и становится источником фактов:
цифры на слайдах сверяются с ним так же, как с текстом, введённым вручную.
"""
from __future__ import annotations

import io
import re
from pathlib import Path

MAX_FILE_CHARS = 20_000      # на файл: README или отчёт целиком, но не дамп репозитория
MAX_TOTAL_CHARS = 30_000     # на бриф: столько планировщик читает без потери скорости

TEXT_EXT = {".txt", ".md", ".markdown", ".rst", ".csv", ".tsv", ".json", ".yaml", ".yml", ".html", ".htm"}
SUPPORTED = TEXT_EXT | {".pdf", ".docx", ".pptx"}


def _clean(text: str) -> str:
    text = re.sub(r"<[^>]+>", " ", text) if "<" in text and ">" in text else text
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)
    return text.strip()


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _pdf(data: bytes) -> str:
    import pymupdf

    from ..core.mupdf_lock import LOCK
    with LOCK:
        doc = pymupdf.open(stream=data, filetype="pdf")
        try:
            return "\n".join(p.get_text() for p in doc)
        finally:
            doc.close()


def _docx(data: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _pptx(data: bytes) -> str:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    parts = []
    for i, s in enumerate(prs.slides, 1):
        texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        if s.has_notes_slide and s.notes_slide.notes_text_frame.text.strip():
            texts.append("Заметки: " + s.notes_slide.notes_text_frame.text)
        if texts:
            parts.append(f"Слайд {i}: " + " / ".join(texts))
    return "\n".join(parts)


def extract(name: str, data: bytes) -> str:
    ext = Path(name).suffix.lower()
    if ext not in SUPPORTED:
        raise ValueError(f"формат {ext or 'без расширения'} не поддерживается — нужен txt, md, pdf, docx или pptx")
    if ext == ".pdf":
        text = _pdf(data)
    elif ext == ".docx":
        text = _docx(data)
    elif ext == ".pptx":
        text = _pptx(data)
    else:
        text = _decode(data)
    text = _clean(text)
    if not text:
        raise ValueError("в файле нет текста (скан или пустой документ)")
    return text[:MAX_FILE_CHARS]


def combine(files: list[tuple[str, str]]) -> str:
    """Склеивает материалы с заголовками файлов, укладываясь в общий лимит."""
    out, left = [], MAX_TOTAL_CHARS
    for name, text in files:
        if left <= 0:
            break
        chunk = text[:left]
        out.append(f"=== {name} ===\n{chunk}")
        left -= len(chunk)
    return "\n\n".join(out)
