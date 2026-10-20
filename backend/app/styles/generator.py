"""Генератор .pptx-пака встроенного стиля.

Пак — обычный шаблон из семи слайдов (обложка, план, раздел, три контентных образца, финал). Он проходит через
тот же парсер, что и корпоративные шаблоны, поэтому весь пайплайн (рецепты, аудит, исправления, экспорт)
работает со стилями без изменений. Виды служебных слайдов задаются метками {{cover.title}} и т. п.

Сетка — 960×540 «пикселей» (1 px = 12700 EMU = 1 pt), как у Kimi: поля 64 px, кикер, заголовок, тонкая линия,
подвал с номером страницы.
"""
from __future__ import annotations

import threading
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from .model import StyleSpec, _mix

PX = 12700
W, H = 960, 540
M = 64                      # поля
CONTENT_W = W - 2 * M


_GRID_LOCK = threading.Lock()      # поля пака задаются брендбуком: сетка — общая для модуля, сборки по очереди


def px(v: float) -> Emu:
    return Emu(int(round(v * PX)))


class _Slide:
    def __init__(self, prs: Presentation, spec: StyleSpec, dark: bool, split: bool = False):
        self.spec = spec
        self.p = spec.palette
        self.dark = dark
        self.s = prs.slides.add_slide(prs.slide_layouts[6])   # Blank
        bg = self.s.background.fill
        bg.solid()
        bg.fore_color.rgb = RGBColor.from_string(self.p.dark_bg if dark and not split else self.p.bg)
        if split:
            # цветной блок на левые 58%: заголовок на нём, мотив — на светлой части
            self.rect("split-block", 0, 0, W * 0.58, H, fill=self.p.dark_bg)

    # ------------------------------------------------------------ примитивы
    def rect(self, name: str, x, y, w, h, fill: str | None = None, line: str | None = None, line_w: float = 0.75,
             shape=MSO_SHAPE.RECTANGLE, radius: float | None = None):
        if radius:
            shape = MSO_SHAPE.ROUNDED_RECTANGLE
        sh = self.s.shapes.add_shape(shape, px(x), px(y), px(w), px(h))
        if radius:
            sh.adjustments[0] = max(0.0, min(0.5, radius / max(1.0, min(w, h))))
        sh.name = name
        if fill:
            sh.fill.solid()
            sh.fill.fore_color.rgb = RGBColor.from_string(fill)
        else:
            sh.fill.background()
        if line:
            sh.line.color.rgb = RGBColor.from_string(line)
            sh.line.width = Pt(line_w)
        else:
            sh.line.fill.background()
        # без p:style — иначе тема python-pptx добавляет фигурам тень и обводку
        st = sh._element.find(qn("p:style"))
        if st is not None:
            sh._element.remove(st)
        sh.text_frame.text = ""
        return sh

    def text(self, name: str, x, y, w, h, text: str, size: float, color: str, font: str | None = None,
             bold: bool = False, caps: bool = False, spacing: float | None = None, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, line: float = 1.0):
        tb = self.s.shapes.add_textbox(px(x), px(y), px(w), px(h))
        tb.name = name
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = anchor
        face = font or self.spec.fonts.body
        for i, para in enumerate(text.split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = line
            r = p.add_run()
            r.text = para
            f = r.font
            f.size = Pt(size)
            f.bold = bold
            f.name = face
            f.color.rgb = RGBColor.from_string(color)
            rpr = r._r.get_or_add_rPr()
            if caps:
                rpr.set("cap", "all")
            if spacing:
                rpr.set("spc", str(int(spacing * 100)))
            for tag in ("a:ea", "a:cs"):       # кириллица и прочие письменности — тем же шрифтом
                el = rpr.find(qn(tag))
                if el is None:
                    el = etree.SubElement(rpr, qn(tag))
                el.set("typeface", face)
        return tb

    # ------------------------------------------------------------ хром
    def kicker(self, text: str, y: float = 40, color: str | None = None):
        c = self.spec.chrome
        col = color or (self.p.accent if self.dark else self.p.primary)
        return self.text("kicker", M, y, CONTENT_W, 16, text, 10.5, col, bold=True, caps=c.kicker_caps,
                         spacing=c.kicker_spacing)

    def title(self, text: str, y: float = 60, h: float = 64, size: float | None = None, w: float | None = None):
        c = self.spec.chrome
        w = CONTENT_W if w is None else w           # поля пака задаются при сборке (брендбук)
        col = self.p.dark_text if self.dark else self.p.text
        return self.text("title", M, y, w, h, text, size or c.title_size, col, font=self.spec.fonts.heading,
                         bold=c.title_bold, line=0.95)

    def logo(self, x: float, y: float, h: float, right: bool = False) -> None:
        """Логотип брендбука: на тёмном слайде — светлый вариант (если есть), высота h, ширина — по пропорциям."""
        path = (self.spec.logo_dark or self.spec.logo) if self.dark else (self.spec.logo or self.spec.logo_dark)
        # правила брендбука: не меньше минимальной высоты; охранное поле — отступ от края слайда
        if self.spec.logo_min_h:
            h = max(h, self.spec.logo_min_h)
        if self.spec.logo_clear:
            pad = self.spec.logo_clear * h
            y = max(y, pad)
            x = min(x, W - pad) if right else max(x, pad)
        if not path or not Path(path).exists():
            return
        from PIL import Image
        with Image.open(path) as im:
            w = h * im.width / max(1, im.height)
        w = min(w, 220)
        hh = w * im.height / max(1, im.width)
        pic = self.s.shapes.add_picture(path, px(x - w if right else x), px(y), px(w), px(hh))
        pic.name = "logo"

    def rule(self, y: float = 132):          # ниже двухстрочного заголовка (28 pt × 2 строки)
        if self.spec.chrome.rule:
            self.rect("rule", M, y, CONTENT_W, 1, fill=self.p.rule)

    def footer(self, deck: str, number: int):
        c = self.spec.chrome
        col = _mix(self.p.dark_text, self.p.dark_bg, 0.45) if self.dark else self.p.muted
        if c.footer:
            self.text("footer", M, 506, 600, 14, deck, 9, col)
        if c.pagenum:
            self.text("pagenum", W - M - 60, 506, 60, 14, str(number), 9, col, align=PP_ALIGN.RIGHT)

    def motif(self, x: float, y: float, s: float = 170, light: bool = False):
        """Декоративный мотив обложки и финала — одна геометрическая идея. light — на светлом фоне."""
        m = self.spec.chrome.motif
        p = self.p
        soft = _mix(p.bg, p.primary, 0.1) if light else _mix(p.dark_bg, p.dark_text, 0.08)
        strong = p.primary if light else p.accent
        if m == "squares":
            self.rect("motif-a", x, y, s, s, fill=soft)
            self.rect("motif-b", x + s * 0.36, y + s * 0.36, s, s, fill=strong if light else p.accent)
        elif m == "bars":
            for i, k in enumerate((1.0, 0.72, 0.46)):
                self.rect(f"motif-{i}", x + i * s * 0.34, y + s * (1 - k), s * 0.24, s * k + s * 0.36,
                          fill=p.accent if i == 0 else soft)
        elif m == "circle":
            self.rect("motif-a", x, y, s * 1.2, s * 1.2, fill=soft, shape=MSO_SHAPE.OVAL)
            self.rect("motif-b", x + s * 0.7, y + s * 0.7, s * 0.5, s * 0.5, fill=p.accent, shape=MSO_SHAPE.OVAL)
        elif m == "frame":
            self.rect("motif-a", x, y, s * 1.3, s * 1.1, line=p.accent, line_w=2)
            self.rect("motif-b", x + s * 0.3, y + s * 0.3, s * 1.3, s * 1.1, fill=soft)
        elif m == "none" and light:
            self.rect("motif-a", x + s * 0.9, y, 4, s * 1.6, fill=p.accent)


def build_pack(spec: StyleSpec, out: Path) -> Path:
    """Собирает пак стиля в out (.pptx) и возвращает путь. Поля — из брендбука (spacing.margin), иначе 64 px."""
    global M, CONTENT_W
    with _GRID_LOCK:
        M = int(max(40, min(100, spec.margin))) if spec.margin else 64
        CONTENT_W = W - 2 * M
        try:
            return _build_pack(spec, out)
        finally:
            M, CONTENT_W = 64, W - 2 * 64


def _build_pack(spec: StyleSpec, out: Path) -> Path:
    p, f, c = spec.palette, spec.fonts, spec.chrome
    prs = Presentation()
    prs.slide_width, prs.slide_height = px(W), px(H)
    deck = "Название презентации · Автор"
    dark_sec = c.dark_sections
    cover_dark = c.cover in ("dark", "split")
    split = c.cover == "split"
    fg = p.dark_text if cover_dark else p.text
    fg_soft = _mix(p.dark_text, p.dark_bg, 0.18) if cover_dark else p.muted
    fg_meta = _mix(p.dark_text, p.dark_bg, 0.35) if cover_dark else p.muted
    tw = 470 if split else 600
    k_split = 0.86 if split else 1.0

    def cover_motif(sl: _Slide, y: float, s: float):
        if split:
            sl.motif(640, y, s, light=True)
        else:
            sl.motif(700, y, s, light=not cover_dark)

    # 1. Обложка
    s = _Slide(prs, spec, dark=cover_dark, split=split)
    s.kicker("Организация · подразделение", y=64, color=p.accent if cover_dark else p.primary)
    s.logo(W - M, 52, 30, right=True)
    s.text("cover-title", M, 168, tw, 150, "{{cover.title}}", c.cover_title_size * k_split, fg, font=f.heading,
           bold=c.title_bold, line=0.95, anchor=MSO_ANCHOR.BOTTOM)
    s.rect("accent-line", M, 336, 56, 4, fill=p.accent)
    s.text("cover-subtitle", M, 356, tw, 50, "Подзаголовок: о чём и для кого эта презентация", 18, fg_soft)
    # метаданные — выше нижней полосы, иначе парсер считает их хромом
    s.text("cover-meta", M, 420, tw, 18, "Автор · дата", 11, fg_meta)
    cover_motif(s, 150, 170)

    # 2. План
    s = _Slide(prs, spec, dark=False)
    s.kicker("Содержание")
    s.title("{{agenda.title}}")
    s.rule()
    for i, (h, t) in enumerate([("Контекст и задача", "С чего начинаем и почему это важно"),
                                ("Решение", "Что предлагаем и как это работает"),
                                ("Результаты", "Что получилось и чем подтверждается"),
                                ("Следующие шаги", "Что нужно сделать дальше")]):
        y = 160 + i * 80
        s.text(f"agenda-num-{i}", M, y, 60, 40, f"0{i + 1}", 26, p.primary, font=f.heading, bold=True)
        s.text(f"agenda-head-{i}", M + 80, y + 2, 360, 24, h, 16, p.text, bold=True)
        s.text(f"agenda-text-{i}", M + 460, y + 4, 372, 40, t, 12.5, p.muted)
        if i < 3:
            s.rect(f"agenda-sep-{i}", M, y + 66, CONTENT_W, 1, fill=p.rule)
    s.footer(deck, 2)
    s.logo(W - M, 36, 20, right=True)

    # 3. Раздел
    s = _Slide(prs, spec, dark=dark_sec)
    s.kicker("Раздел 01", y=150, color=p.accent)
    s.text("section-title", M, 176, 620, 120, "{{section.title}}", 40, p.dark_text if dark_sec else p.text,
           font=f.heading, bold=c.title_bold, line=0.95)
    s.rect("accent-line", M, 318, 56, 4, fill=p.accent)
    s.text("section-sub", M, 340, 620, 40, "Что будет в этом разделе", 15,
           _mix(p.dark_text, p.dark_bg, 0.2) if dark_sec else p.muted)
    s.text("section-num", W - M - 340, 310, 340, 180, "01", 130, _mix(p.dark_bg, p.dark_text, 0.12) if dark_sec
           else _mix(p.bg, p.primary, 0.12), font=f.heading, bold=True, align=PP_ALIGN.RIGHT)
    s.footer(deck, 3)

    # 4. Контент: нумерованные строки
    s = _Slide(prs, spec, dark=False)
    s.kicker("Тема · подтема")
    s.title("Заголовок слайда — вывод, а не тема")
    s.rule()
    s.text("lead", M, 150, CONTENT_W, 40, "Вводная фраза: одно предложение, которое объясняет, зачем этот слайд.",
           14, p.text, line=1.2)
    for i in range(3):
        y = 208 + i * 90
        s.text(f"row-num-{i}", M, y, 48, 36, f"0{i + 1}", 22, p.accent if i == 0 else p.primary, font=f.heading,
               bold=True)
        s.text(f"row-head-{i}", M + 64, y + 2, 300, 22, "Короткий тезис", 15, p.text, bold=True)
        s.text(f"row-text-{i}", M + 64, y + 28, 760, 40,
               "Пояснение в одну-две строки: что именно происходит и почему это важно для аудитории.", 12.5,
               p.muted, line=1.2)
    s.footer(deck, 4)
    s.logo(W - M, 36, 20, right=True)

    # 5. Контент: две колонки с разделителем
    s = _Slide(prs, spec, dark=False)
    s.kicker("Тема · сравнение")
    s.title("Две стороны одного решения")
    s.rule()
    for i, x in enumerate((M, M + CONTENT_W / 2 + 24)):
        s.text(f"col-head-{i}", x, 154, CONTENT_W / 2 - 48, 22, "ПОДЗАГОЛОВОК КОЛОНКИ", 11, p.primary, bold=True,
               caps=True, spacing=1.5)
        s.text(f"col-body-{i}", x, 186, CONTENT_W / 2 - 48, 220,
               "• первый пункт колонки\n• второй пункт колонки\n• третий пункт колонки", 13.5, p.text, line=1.35)
    s.rect("col-divider", M + CONTENT_W / 2, 154, 1, 306, fill=p.rule)
    s.footer(deck, 5)
    s.logo(W - M, 36, 20, right=True)

    # 6. Контент: три карточки (образец стиля карточки)
    s = _Slide(prs, spec, dark=False)
    s.kicker("Тема · ключевые элементы")
    s.title("Три опоры решения")
    s.rule()
    cw = (CONTENT_W - 2 * 24) / 3
    for i in range(3):
        x = M + i * (cw + 24)
        s.rect(f"card-{i}", x, 164, cw, 220, fill=p.surface if c.card_fill else None, line=p.rule, radius=spec.radius)
        s.text(f"card-num-{i}", x + 20, 184, 60, 30, f"0{i + 1}", 20, p.accent if i == 0 else p.primary,
               font=f.heading, bold=True)
        s.text(f"card-head-{i}", x + 20, 226, cw - 40, 24, "Название опоры", 15, p.text, bold=True)
        s.text(f"card-text-{i}", x + 20, 258, cw - 40, 100, "Две-три строки о том, что это и зачем нужно.", 12.5,
               p.muted, line=1.25)
    s.footer(deck, 6)
    s.logo(W - M, 36, 20, right=True)

    # 7. Финал
    s = _Slide(prs, spec, dark=cover_dark, split=split)
    s.kicker("Итоги", y=150, color=p.accent if cover_dark else p.primary)
    s.text("closing-title", M, 176, tw, 120, "{{closing.title}}", 40 * k_split, fg, font=f.heading,
           bold=c.title_bold, line=0.95)
    s.rect("accent-line", M, 318, 56, 4, fill=p.accent)
    s.text("closing-sub", M, 340, tw, 30, "Контакты и следующий шаг", 15, fg_soft)
    s.logo(W - M, 52, 30, right=True)
    # образец списка выводов: даёт итоговому холсту слот body, и выводы попадают на финал, а не в карточки
    s.text("closing-body", M, 384, tw - 40, 110, "Первый вывод презентации\nВторой вывод презентации\nТретий вывод",
           14, _mix(p.dark_text, p.dark_bg, 0.12) if cover_dark else p.text, line=1.3)
    cover_motif(s, 200, 150)

    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out
