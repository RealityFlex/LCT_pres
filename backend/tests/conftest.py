"""Общие фикстуры: изолированный каталог данных и синтетический шаблон, собранный python-pptx.

Реальные шаблоны хакатона лежат в data/ (не в репозитории) — тесты от них не зависят,
а интеграционные проверки на них пропускаются, если файлов нет.
"""
from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path

import pytest

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))

# данные тестов — во временном каталоге, а не в data/ проекта (до первого get_settings())
_TMP = Path(tempfile.mkdtemp(prefix="lekalo-tests-"))
os.environ["DATA_DIR"] = str(_TMP)

from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.util import Emu, Inches, Pt  # noqa: E402

FIXTURES = BACKEND / "tests" / "fixtures"
PRIMARY = "1F4E9C"
ACCENT = "E0457B"
TEXT = "1E1E24"


def _text(slide, x, y, w, h, text, size, color=TEXT, bold=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    r = tf.paragraphs[0].add_run()
    r.text = text
    r.font.name = "Arial"
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    return tb


def _rect(slide, x, y, w, h, fill, geom=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(geom, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = RGBColor.from_string(fill)
    s.line.fill.background()
    return s


def build_synthetic_template(path: Path) -> Path:
    """Пять слайдов: титул, раздел, карточки в подложках, колонки без подложек, финал."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    # 1. титул
    s = prs.slides.add_slide(blank)
    _rect(s, 0, 0, 13.333, 0.35, PRIMARY, MSO_SHAPE.RECTANGLE)
    _text(s, 0.8, 2.6, 10, 1.2, "Название презентации", 44, PRIMARY, True)
    _text(s, 0.8, 3.9, 10, 0.6, "Подзаголовок и дата", 20)
    # 2. раздел (тёмный)
    s = prs.slides.add_slide(prs.slide_layouts[2])
    for ph in list(s.placeholders):
        ph._element.getparent().remove(ph._element)
    _rect(s, 0, 0, 13.333, 7.5, PRIMARY, MSO_SHAPE.RECTANGLE)
    _text(s, 0.8, 3.0, 10, 1.2, "Раздел первый", 48, "FFFFFF", True)
    # 3. карточки в подложках
    s = prs.slides.add_slide(blank)
    _text(s, 0.8, 0.5, 11.5, 0.8, "Преимущества решения", 30, PRIMARY, True)
    for i in range(3):
        x = 0.8 + i * 4.0
        _rect(s, x, 1.8, 3.7, 3.6, "EEF2F8")
        _text(s, x + 0.3, 2.1, 3.1, 0.6, f"Пункт {i + 1}", 20, PRIMARY, True)
        _text(s, x + 0.3, 2.8, 3.1, 2.2, "Короткое пояснение к пункту, которое занимает две-три строки текста.", 14)
    # 4. колонки без подложек (группа по стилю текста)
    s = prs.slides.add_slide(blank)
    _text(s, 0.8, 0.5, 11.5, 0.8, "Ключевые возможности", 30, PRIMARY, True)
    for i in range(3):
        x = 0.8 + i * 4.0
        _text(s, x, 2.0, 3.5, 0.5, f"Возможность {i + 1}", 18, ACCENT, True)
        _text(s, x, 2.6, 3.5, 1.8, "Описание возможности в одно-два предложения обычным текстом.", 14)
    # 5. финал
    s = prs.slides.add_slide(blank)
    _text(s, 0.8, 3.0, 11.5, 1.2, "Спасибо за внимание", 44, PRIMARY, True)
    prs.save(str(path))
    return path


@pytest.fixture(scope="session")
def synthetic_pptx(tmp_path_factory) -> Path:
    return build_synthetic_template(tmp_path_factory.mktemp("tpl") / "synthetic.pptx")


def soffice_available() -> bool:
    from app.config import get_settings
    try:
        return bool(get_settings().soffice)
    except Exception:
        return False


needs_soffice = pytest.mark.skipif(not soffice_available(), reason="LibreOffice не установлен")


@pytest.fixture(scope="session")
def synthetic_profile(synthetic_pptx):
    """Профиль синтетического шаблона (без LLM). Требует LibreOffice для рендера."""
    if not soffice_available():
        pytest.skip("LibreOffice не установлен")
    from app.parsing.profile import parse_template
    return parse_template(synthetic_pptx, name="Синтетический", use_llm=False)
