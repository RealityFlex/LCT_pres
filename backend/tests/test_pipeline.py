"""Интеграционные тесты: разбор шаблона → сборка трёх вариантов на фиксированном плане → рендер → аудит.

Модели не вызываются (контент — фикстура). Нужен LibreOffice; без него тесты пропускаются.
"""
from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation

from tests.conftest import ACCENT, FIXTURES, PRIMARY, needs_soffice

pytestmark = needs_soffice


def test_profile_design_system(synthetic_profile):
    p = synthetic_profile
    kinds = {c.kind for c in p.canvases}
    assert {"title", "section", "content"} <= kinds
    hexes = {(c["hex"] if isinstance(c, dict) else c.hex) for c in p.palette.colors}
    assert PRIMARY in hexes
    assert p.typography.title.size >= 28 and p.typography.title.bold
    assert any(c.pattern and c.pattern.kind == "cards" for c in p.patterns)
    # сетка шаблона: поля 0.8″ слева
    assert abs(p.margins.x / 914400 - 0.8) < 0.15


@pytest.fixture(scope="module")
def built(synthetic_profile, tmp_path_factory):
    from app.config import get_settings
    from app.pipeline import build_and_render, compose_variant
    from app.planning.models import DeckContent
    from app.planning.variants import build_variant
    content = DeckContent.model_validate_json((FIXTURES / "demo_content.json").read_text(encoding="utf-8"))
    out = {}
    for v in get_settings().variants:
        plan = build_variant(content, synthetic_profile, v["id"], v["name"], v["strategy"], v.get("description", ""))
        layouts = compose_variant(synthetic_profile, content, plan)
        vdir = tmp_path_factory.mktemp(f"v{v['id']}")
        files = build_and_render(synthetic_profile, layouts, vdir, f"deck_{v['id']}")
        out[v["id"]] = (plan, layouts, files)
    return out


def test_three_variants_are_built_and_open(built):
    assert set(built) == {"A", "B", "C"}
    counts = {}
    for vid, (plan, layouts, files) in built.items():
        prs = Presentation(str(files["pptx"]))
        assert len(prs.slides) == len(layouts) == len(plan.slides)
        assert Path(files["pdf"]).exists() and len(files["pngs"]) == len(layouts)
        counts[vid] = len(layouts)
    # компактный вариант не длиннее классического
    assert counts["C"] <= counts["A"]


def test_variants_differ_in_layout(built):
    recipes = {vid: [lay.recipe for lay in layouts] for vid, (_, layouts, _) in built.items()}
    assert recipes["A"] != recipes["B"] or recipes["A"] != recipes["C"]


def test_native_objects_not_pictures(built):
    """Диаграммы и таблицы — нативные объекты PPTX, а не картинки."""
    _, _, files = built["A"]
    prs = Presentation(str(files["pptx"]))
    kinds = set()
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_chart:
                kinds.add("chart")
            if sh.has_table:
                kinds.add("table")
    assert kinds, "в колоде по фикстуре должны быть нативные диаграммы или таблицы"


def test_deterministic_audit_has_no_layout_errors(built, synthetic_profile):
    from app.audit.deterministic import DeterministicAuditor
    bad = {"out_of_bounds", "overlap", "text_overflow", "text_cut", "font_not_in_template", "empty_slide"}
    for vid, (_, layouts, files) in built.items():
        issues, checks = DeterministicAuditor(synthetic_profile).run(files["pptx"], layouts, files["pngs"], len(layouts))
        assert "text_overflow" in checks and "contrast" in checks
        errors = [f"{i.check}@{i.slide}: {i.message}" for i in issues if i.severity == "error" and i.check in bad]
        assert not errors, f"вариант {vid}: {errors}"


REAL = sorted((Path(__file__).resolve().parents[2] / "data" / "templates").glob("*/template.pptx"))


@pytest.mark.skipif(not REAL, reason="реальные шаблоны не загружены в data/templates")
@pytest.mark.parametrize("tpl", REAL[:6], ids=lambda p: p.parent.name)
def test_real_templates_give_canvases(tpl):
    """На каждом загруженном шаблоне находятся холсты титула и контента (детерминированно, без LLM)."""
    from app.parsing.analyze import TemplateAnalyzer
    from app.parsing.raster import is_raster
    an = TemplateAnalyzer(tpl)
    slides = an.run()
    if is_raster(an, slides):
        pytest.skip("растровый шаблон — проверяется восстановлением по пикселям")
    usable = [s for s in slides if not s.service]
    assert any(s.title is not None for s in usable)
    assert any(s.kind == "content" for s in usable)
