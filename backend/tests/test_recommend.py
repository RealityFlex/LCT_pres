"""Профиль датасета, аналитическая задача и рекомендатель визуализаций (без модели)."""
import io
from pathlib import Path

import pandas as pd
import pytest
from pptx import Presentation

from app.analytics.analyze import analyze, unreadable_reason
from app.analytics.ingest import Table, load
from app.analytics.recommend import build_chart, goal_for
from app.analytics.story import apply_viz
from app.layout.builder import DeckBuilder
from app.layout.ir import El
from app.parsing.model import Box
from app.planning.models import ContentSlide

DEMO = Path(__file__).resolve().parents[2] / "examples" / "data_demo"


def _table(name):
    tables, _ = load(name, (DEMO / name).read_bytes())
    return tables[0]


def _ins(name, goal=""):
    t = _table(name)
    return analyze(t, goal_for(goal, t.name, [str(c) for c in t.df.columns]))


@pytest.mark.parametrize("name,task,viz", [
    ("sales_monthly.csv", "timeseries", "line"),      # 36 месяцев → линия, столбцы сливаются
    ("channels.json", "shares", "doughnut"),          # 6 долей одного целого → кольцо
    ("regions.xlsx", "comparison", "bar"),            # 9 регионов с длинными названиями → полосы
    ("checks.csv", "distribution", "histogram"),
    ("report_q3.txt", "timeseries", "clustered"),     # 7 кварталов, два ряда → сгруппированные столбцы
])
def test_task_and_first_option(name, task, viz):
    ins = _ins(name)
    assert ins.nature == task and ins.viz == viz
    assert 2 <= len(ins.options) <= 3 and all(o["why"] for o in ins.options)
    assert [o["score"] for o in ins.options] == sorted((o["score"] for o in ins.options), reverse=True)


def test_rejected_explained():
    ins = _ins("regions.xlsx")
    assert any("круговая" in r and "6" in r for r in ins.options[0]["rejected"])


def test_goal_switches_task_to_relationship():
    plain = _ins("products.xlsx")
    goal = _ins("products.xlsx", "Показать, как цена влияет на продажи.")
    assert plain.nature == "reference" and goal.nature == "relationship" and goal.viz == "scatter"
    assert goal.chart["points"]["x"] and len(goal.chart["trend"]) == 2
    assert goal.chart["x_title"].startswith("Цена") and goal.chart["y_title"].startswith("Продажи")
    assert any("Внутри групп" in f for f in goal.facts)


def test_goal_applies_only_to_its_dataset():
    # предложение про цену и продажи не касается выручки по регионам
    t = _table("regions.xlsx")
    assert goal_for("Показать, как цена влияет на продажи.", t.name, [str(c) for c in t.df.columns]) == ""


def test_goal_compare_changes_gaps_dataset():
    plain = _ins("survey.csv")
    goal = _ins("survey.csv", "Сравнить оценку NPS по городам.")
    assert plain.nature == "gaps" and goal.nature == "comparison"


def test_same_type_regardless_of_file_format():
    t = _table("regions.xlsx")
    buf = io.StringIO()
    t.df.to_csv(buf, sep=";", index=False)
    csv, _ = load("regions.csv", buf.getvalue().encode("utf-8"))
    a, b = analyze(t), analyze(csv[0])
    assert (a.nature, a.viz, [o["type"] for o in a.options]) == (b.nature, b.viz, [o["type"] for o in b.options])


def test_profile_roles_hierarchy_outliers():
    p = _ins("regions.xlsx").profile
    assert ["Регион", "Федеральный округ"] in p["hierarchies"]
    assert "Регион" in p["dimensions"] and "Выручка 2025, млн ₽" in p["metrics"]
    checks = _ins("checks.csv").profile
    col = next(c for c in checks["columns"] if c["name"].startswith("Сумма"))
    assert col["role"] == "metric" and col["outliers"] > 0
    assert checks["time_axis"] == "Время" and "выбросы" in checks["summary"]


def test_unreadable_dataset_reason():
    t = Table(name="notes", df=pd.DataFrame({"Комментарий": [f"текст {i} про что-то своё" for i in range(30)]}), source="csv")
    assert unreadable_reason(t) and analyze(t) is None


def test_every_option_builds_pptx():
    """Любой вариант любого датасета превращается в слайд и собирается в PPTX без ошибок."""
    prs = Presentation()
    b = object.__new__(DeckBuilder)          # сборщику диаграмм профиль шаблона не нужен
    for name in ("sales_monthly.csv", "channels.json", "regions.xlsx", "checks.csv", "survey.csv"):
        ins = _ins(name)
        for o in ins.options:
            s = ContentSlide(id="s", title="t")
            apply_viz(s, ins, o["type"])
            assert s.intent in ("chart", "table", "stats")
            if s.chart is None:
                continue
            ch = s.chart.model_dump()
            ch["colors"] = ["0077FF", "FF7700", "22AA55", "AA22AA", "888888", "333333"]
            sl = prs.slides.add_slide(prs.slide_layouts[6])
            b._chart(sl, El(id="c", kind="chart", box=Box(x=0, y=0, w=6000000, h=3500000), chart=ch))
    ins = _ins("products.xlsx", "Показать, как цена влияет на продажи.")
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    ch = build_chart(ins, "scatter")
    b._chart(sl, El(id="c", kind="chart", box=Box(x=0, y=0, w=6000000, h=3500000), chart=ch))
    prs.save(io.BytesIO())


def test_many_categories_split_and_failed_slide():
    from app.analytics.story import failed_slides, tail_slide
    from app.parsing.model import TemplateProfile
    # 40 городов с длинным хвостом: за пределами первых 10 — больше четверти итога
    df = pd.DataFrame({"Город": [f"Город {i:02d}" for i in range(40)], "Продажи, шт": [1000 - 20 * i for i in range(40)]})
    ins = analyze(Table(name="cities", df=df, source="csv"))
    assert ins.nature == "comparison" and ins.agg["tail_share"] > 0.25
    first = ContentSlide(id="d1", title="t")
    apply_viz(first, ins, ins.viz)
    extra = tail_slide(ins, first)
    assert extra is not None and len(first.chart.categories) == 10
    # у полос первая категория рисуется снизу: «Прочие» — внизу, 11-й город — сверху
    assert extra.chart.categories[0] == "Прочие" and extra.chart.categories[-1] == "Город 10"
    assert len(extra.chart.categories) == 21
    sl = failed_slides([{"file": "notes.csv", "name": "notes", "reason": "нет чисел"}], TemplateProfile.model_construct())
    assert sl and "не вошёл в анализ" in sl[0].title and sl[0].items[0].text == "notes.csv: нет чисел"


def test_data_deck_skeleton_offline():
    """Без модели каркас всё равно цельный: главное → план → данные → выводы → финал (методика — в заметках), числа — из фактов."""
    import asyncio

    from app.analytics.story import data_deck, load_insights
    from app.parsing.model import Narrative, TemplateProfile
    from app.planning.models import Brief

    class NoRun:
        async def call(self, *a, **k):
            raise RuntimeError("offline")

    paths = [p for p in sorted(DEMO.iterdir()) if p.suffix != ".py" and p.is_file()]
    failed, cross = [], []
    ins, _ = load_insights(paths, failed=failed, cross=cross)
    prof = TemplateProfile.model_construct(narrative=Narrative(), title_chars=60)
    deck = asyncio.run(data_deck(ins, failed, cross, Brief(topic="Зерно"), prof, NoRun()))
    intents = [s.intent for s in deck.slides]
    assert intents[:3] == ["title", "stats", "agenda"] and intents[-1] == "closing"
    assert "table" not in intents[-3:] and any("Методика:" in s.notes for s in deck.slides)
    assert sum(1 for s in deck.slides if s.data.get("options")) == len(ins) == 7
    assert all(s.source_note.startswith("Источник:") for s in deck.slides if s.data)
    assert len(deck.slides[1].items) == 3 and cross


def test_stacked_from_parts_and_pie_options():
    reg = _ins("regions.xlsx")
    assert "stacked_bar" in [o["type"] for o in reg.options]
    ch = build_chart(reg, "stacked_bar")
    assert [s["name"] for s in ch["series"]][0].startswith("Формат") and ch["categories"][-1] == "Москва"
    assert any(f.startswith("Состав итога") for f in reg.facts)
    assert "pie" in [o["type"] for o in _ins("channels.json").options]
